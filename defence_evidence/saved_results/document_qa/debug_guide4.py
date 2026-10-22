import importlib.util, traceback
from docx import Document
p=r"C:\Users\HP\Desktop\KneeAssist-AI\scripts\build_guide_aligned_dissertation.py"
s=importlib.util.spec_from_file_location("b",p); b=importlib.util.module_from_spec(s); s.loader.exec_module(b)
origmove=b.move_heading_block_before
def m(doc,a,c):
 print('before move chapter2',any(x.text.strip()=='CHAPTER 2: LITERATURE REVIEW' for x in doc.paragraphs)); r=origmove(doc,a,c); print('after move chapter2',any(x.text.strip()=='CHAPTER 2: LITERATURE REVIEW' for x in doc.paragraphs)); return r
b.move_heading_block_before=m
orig=b.reorder_h2_blocks
def wrap(doc,title,desired):
 print('reorder',title,'target exists',any(x.text.strip()==title for x in doc.paragraphs))
 return orig(doc,title,desired)
b.reorder_h2_blocks=wrap
d=Document(b.SOURCE)
try: b.guide_content(d)
except Exception: traceback.print_exc()
print('final target',[(x.text,x.style.name) for x in d.paragraphs if 'CHAPTER 2' in x.text])
