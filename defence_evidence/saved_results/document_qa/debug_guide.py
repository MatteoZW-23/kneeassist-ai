import importlib.util
from docx import Document
p=r"C:\Users\HP\Desktop\KneeAssist-AI\scripts\build_guide_aligned_dissertation.py"
s=importlib.util.spec_from_file_location("b",p); b=importlib.util.module_from_spec(s); s.loader.exec_module(b)
d=Document(b.SOURCE)
b.rename(d,{"1.1 Background":"1.1 Introduction","1.2 Problem statement":"1.3 Problem statement","1.3 Aim":"1.4 Project aim","1.4 Proposed revised objectives":"1.5 Project objectives","1.6 Scope and boundaries":"1.8 Project scope","1.7 Motivation and significance":"1.6 Motivation of the project","1.8 Stakeholders":"1.7 Stakeholders","1.9 Assumptions":"1.9 Project assumptions","1.11 Chapter summary":"1.11 Chapter summary"})
print('before', [p.text for p in d.paragraphs if p.text=='CHAPTER 2: LITERATURE REVIEW'])
b.delete_block(d,"1.5 Research questions")
print('after delete', [p.text for p in d.paragraphs if p.text=='CHAPTER 2: LITERATURE REVIEW'])
b.remove_table_with_headers(d,{"Question","Project evidence"})
print('after table', [p.text for p in d.paragraphs if p.text=='CHAPTER 2: LITERATURE REVIEW'])
print('remaining h1',[(p.text,p.style.name) for p in d.paragraphs if p.style.name=='Heading 1'][:10])
