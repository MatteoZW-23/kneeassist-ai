import importlib.util
from docx import Document
p=r"C:\Users\HP\Desktop\KneeAssist-AI\scripts\build_guide_aligned_dissertation.py"
s=importlib.util.spec_from_file_location("b",p); b=importlib.util.module_from_spec(s); s.loader.exec_module(b)
d=Document(b.SOURCE)
b.guide_content(d)
print('success',len(d.paragraphs))
print([p.text for p in d.paragraphs if p.style.name=='Heading 2' and p.text.startswith(('1.','2.'))][:30])
