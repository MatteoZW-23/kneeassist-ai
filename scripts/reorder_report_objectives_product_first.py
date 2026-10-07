"""Put the product-oriented KneeAssist AI objective first throughout the final report."""

from pathlib import Path

from docx import Document


REPORT = Path(r"C:\Users\HP\Desktop\KneeAssist-AI\docs\academic\submission\Mabira_Mathew_R234371Y.docx")


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def find(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Missing paragraph beginning: {prefix}")


def fill_row(table, row_index: int, values: list[str]) -> None:
    for column, value in enumerate(values):
        set_text(table.cell(row_index, column).paragraphs[0], value)


def reorder_chapter_blocks(document: Document, blocks: list, final_boundary) -> None:
    """Reorder the five adjacent Chapter 4 blocks without losing tables or figures."""
    body = document.element.body
    children = list(body)
    starts = [children.index(block._p) for block in blocks]
    end = children.index(final_boundary._p)
    groups = []
    for left, right in zip(starts, starts[1:] + [end]):
        groups.append(children[left:right])
    first = starts[0]
    for element in children[first:end]:
        body.remove(element)
    # Existing order: preprocessing, transfer model, comparison, evaluation, prototype.
    # Required order: prototype, preprocessing, transfer model, comparison, evaluation.
    order = [groups[4], groups[0], groups[1], groups[2], groups[3]]
    location = first
    for group in order:
        for element in group:
            body.insert(location, element)
            location += 1


def main() -> None:
    document = Document(REPORT)
    objectives = [
        "1. To develop a Streamlit-based KneeAssist AI prototype that displays predictions, MRI slices, Grad-CAM attention and an exportable research summary.",
        "2. To prepare and preprocess knee MRI studies for deep-learning classification using normalisation, slice sampling and data-quality checks.",
        "3. To develop a transfer-learning model for automated study-level classification of normal and abnormal knee MRI studies.",
        "4. To compare ResNet-18 and EfficientNet-B0 models and select the better model using validation AUROC.",
        "5. To evaluate the selected model against available MRI reference labels using accuracy, sensitivity, specificity, precision, F1-score and AUROC.",
    ]
    first = find(document, "1. To prepare and preprocess")
    paragraphs = document.paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p._p is first._p)
    for offset, text in enumerate(objectives):
        set_text(paragraphs[index + offset], text)

    # Research-question/objective alignment table and achieved-objective table.
    mapping = [
        ("1", "Can a local prototype present research scores and model attention for supported MRI studies?", "Upload, analysis, visualisation, export and clear-case verification"),
        ("2", "How can knee MRI studies be prepared consistently for deep-learning classification?", "Normalisation, slice sampling, plane checks and data-quality controls"),
        ("3", "Can a transfer-learning model classify knee MRI studies as normal or abnormal?", "Study-level multi-plane classifier and supported predictions"),
        ("4", "Which compared model gives the stronger validation AUROC?", "ResNet-18 and EfficientNet-B0 validation comparison and selected checkpoint"),
        ("5", "What performance does the selected model achieve against available MRI reference labels?", "Per-target metrics, calibration and documented evaluation limitations"),
    ]
    for row, values in enumerate(mapping, 1):
        fill_row(document.tables[3], row, list(values))
    achievements = [
        ("1", "Verified local MRI workflow, score display, Grad-CAM display and exports", "Research usability only; no clinical field study"),
        ("2", "Prepared MRI studies with controlled normalisation, slice sampling and input checks", "MRNet-only training data; guarded support for new research inputs"),
        ("3", "Developed a study-level EfficientNet-B0 transfer-learning classifier", "MRNet-only training; general abnormality, ACL and meniscus outputs"),
        ("4", "Compared ResNet-18 and EfficientNet-B0 using validation AUROC", "EfficientNet-B0 retained as the selected model"),
        ("5", "Reported AUROC, F1, sensitivity, specificity, precision and accuracy", "Available reference labels only; no radiologist-reader study"),
    ]
    for row, values in enumerate(achievements, 1):
        fill_row(document.tables[14], row, list(values))

    # Rename then move the evidence blocks as a single unit, retaining their tables and figures.
    prep = find(document, "4.2.1 MRI preparation")
    model = find(document, "4.2.2 Transfer-learning")
    comparison = find(document, "4.2.3 Model comparison")
    evaluation = find(document, "4.2.4 Evaluation against")
    prototype = find(document, "4.2.5 KneeAssist AI decision-support")
    boundary = find(document, "4.3 Descriptive statistics")
    set_text(prototype, "4.2.1 KneeAssist AI decision-support prototype (Objective 1)")
    set_text(prep, "4.2.2 MRI preparation and preprocessing (Objective 2)")
    set_text(model, "4.2.3 Transfer-learning classification model (Objective 3)")
    set_text(comparison, "4.2.4 Model comparison and selection (Objective 4)")
    set_text(evaluation, "4.2.5 Evaluation against available MRI reference labels (Objective 5)")
    reorder_chapter_blocks(document, [prep, model, comparison, evaluation, prototype], boundary)

    abstract = find(document, "Knee magnetic resonance imaging classification requires")
    set_text(
        abstract,
        abstract.text.replace(
            "The five objectives address MRI preprocessing, transfer-learning classification, model comparison, quantitative evaluation against available reference labels, and a decision-support prototype.",
            "The five objectives address a decision-support prototype, MRI preprocessing, transfer-learning classification, model comparison, and quantitative evaluation against available reference labels.",
        ),
    )
    summary = find(document, "The five objectives were addressed with different")
    set_text(
        summary,
        "The five objectives were addressed with different types of evidence. The decision-support prototype, controlled MRI preprocessing, transfer-learning classification, validation-based model selection and quantitative metrics against available labels are observable technical achievements. The absence of a radiologist-reader study and weak external transfer are findings to report, not reasons to hide the results.",
    )
    document.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
