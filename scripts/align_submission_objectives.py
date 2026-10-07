"""Align the final report with the student's five approved objectives."""

from pathlib import Path

from docx import Document


REPORT = Path(r"C:\Users\HP\Desktop\KneeAssist-AI\docs\academic\submission\Mabira_Mathew_R234371Y.docx")


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def find_paragraph(document: Document, starts_with: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(starts_with):
            return paragraph
    raise RuntimeError(f"Could not find paragraph beginning: {starts_with}")


def paragraph_index(document: Document, paragraph) -> int:
    for index, candidate in enumerate(document.paragraphs):
        if candidate._p is paragraph._p:
            return index
    raise RuntimeError("Paragraph is not in the document body.")


def fill_row(table, row_index: int, values: list[str]) -> None:
    for column, value in enumerate(values):
        set_text(table.cell(row_index, column).paragraphs[0], value)


def reorder_blocks(document: Document, first, boundaries: list) -> None:
    """Move five adjacent Chapter 4 blocks while retaining contained tables."""
    body = document.element.body
    children = list(body)
    first_index = children.index(first._p)
    boundary_indexes = [children.index(item._p) for item in boundaries]
    all_blocks = [
        children[left:right]
        for left, right in zip([first_index] + boundary_indexes[:-1], boundary_indexes)
    ]
    # Current order is model, comparison, evaluation, preprocessing, interface.
    # The achieved-objective order is preprocessing, model, comparison, evaluation, interface.
    approved_order = [all_blocks[3], all_blocks[0], all_blocks[1], all_blocks[2], all_blocks[4]]
    for element in children[first_index:boundary_indexes[-1]]:
        body.remove(element)
    insert_at = first_index
    for block in approved_order:
        for element in block:
            body.insert(insert_at, element)
            insert_at += 1


def main() -> None:
    document = Document(REPORT)
    objectives = [
        "1. To prepare and preprocess knee MRI studies for deep-learning classification using normalisation, slice sampling and data-quality checks.",
        "2. To develop a transfer-learning model for automated study-level classification of normal and abnormal knee MRI studies.",
        "3. To compare ResNet-18 and EfficientNet-B0 models and select the better model using validation AUROC.",
        "4. To evaluate the selected model against available MRI reference labels using accuracy, sensitivity, specificity, precision, F1-score and AUROC.",
        "5. To develop a Streamlit-based KneeAssist AI prototype that displays predictions, MRI slices, Grad-CAM attention and an exportable research summary.",
    ]
    first_objective = find_paragraph(document, "1. To develop an AI model")
    paragraphs = document.paragraphs
    first_index = paragraph_index(document, first_objective)
    for index, objective in enumerate(objectives):
        set_text(paragraphs[first_index + index], objective)

    set_text(
        find_paragraph(document, "Knee magnetic resonance imaging classification requires"),
        "Knee magnetic resonance imaging classification requires more than a model that produces plausible scores: "
        "data separation, reference labels and evaluation conditions determine what those scores can support. "
        "This project developed KneeAssist AI, a local research application for study-level classification of normal "
        "and abnormal knee MRI studies. ACL tear and meniscal tear scores are retained as secondary research outputs. "
        "The five objectives address MRI preprocessing, transfer-learning classification, model comparison, quantitative "
        "evaluation against available reference labels, and a decision-support prototype. A retrospective computational design "
        "used MRNet as the sole training source, with 804 fitting, 226 internal tuning, 100 calibration and 120 official "
        "validation studies. ImageNet-pretrained ResNet-18 and EfficientNet-B0 encoders processed sampled slices from available "
        "imaging planes, followed by study-level aggregation and a multi-label classification head. The selected checkpoint was "
        "the frozen-encoder EfficientNet-B0 warm-up model; attempted fine-tuning and a separate five-family benchmark did not "
        "exceed its internal selection performance. The primary normal/abnormal output achieved AUROC of 0.901 on the official "
        "MRNet validation cohort. The interface supports prepared MRI arrays, cautious NIfTI input and guarded DICOM routes. "
        "Independent inference, Grad-CAM attention, MRI upload, export controls and case clearing were verified; release evidence "
        "recorded 33 passing tests and nine executed notebooks. A separate RSNA Kaggle pilot was not merged because only 58 studies "
        "had complete explicit labels. Patient linkage was unavailable, evaluation material had been reused, and external reference "
        "evidence was limited. No radiologist-reader comparison was conducted. The contribution is an auditable working research "
        "prototype with transparent limitations and reproducible artifacts, rather than a clinically validated diagnostic system. "
        "Independent patient-level evaluation and controlled robustness studies are recommended before stronger claims."
    )

    research_question_rows = [
        ("1", "How can knee MRI studies be prepared consistently for deep-learning classification?", "Normalisation, slice sampling, plane checks and data-quality controls"),
        ("2", "Can a transfer-learning model classify knee MRI studies as normal or abnormal?", "Study-level multi-plane classifier and supported predictions"),
        ("3", "Which compared model gives the stronger validation AUROC?", "ResNet-18 and EfficientNet-B0 validation comparison and selected checkpoint"),
        ("4", "What performance does the selected model achieve against available MRI reference labels?", "Per-target metrics, calibration and documented evaluation limitations"),
        ("5", "Can a local prototype present predictions and model attention for research review?", "Upload, analysis, visualisation, export and clear verification"),
    ]
    for row_index, values in enumerate(research_question_rows, start=1):
        fill_row(document.tables[3], row_index, list(values))

    achievement_rows = [
        ("1", "Prepared MRI studies with controlled normalisation, slice sampling and input checks", "MRNet-only training data; guarded support for new research inputs"),
        ("2", "Developed a study-level EfficientNet-B0 transfer-learning classifier", "MRNet-only training; general abnormality, ACL and meniscus outputs"),
        ("3", "Compared ResNet-18 and EfficientNet-B0 using validation AUROC", "EfficientNet-B0 retained as the selected model"),
        ("4", "Reported AUROC, F1, sensitivity, specificity, precision and accuracy", "Available reference labels only; no radiologist-reader study"),
        ("5", "Verified local MRI workflow, Grad-CAM display and exports", "Research usability only; no clinical field study"),
    ]
    for row_index, values in enumerate(achievement_rows, start=1):
        fill_row(document.tables[14], row_index, list(values))

    # Rename before moving the corresponding XML blocks.
    model = find_paragraph(document, "4.2.1 Automated knee MRI classification")
    features = find_paragraph(document, "4.2.2 MRI feature attention")
    evaluation = find_paragraph(document, "4.2.3 Quantitative model evaluation")
    data = find_paragraph(document, "4.2.4 Comparison with established")
    interface = find_paragraph(document, "4.2.5 KneeAssist AI decision-support")
    next_section = find_paragraph(document, "4.3 Descriptive statistics")
    set_text(data, "4.2.1 MRI preparation and preprocessing (Objective 1)")
    set_text(model, "4.2.2 Transfer-learning classification model (Objective 2)")
    set_text(features, "4.2.3 Model comparison and selection (Objective 3)")
    set_text(evaluation, "4.2.4 Evaluation against available MRI reference labels (Objective 4)")
    set_text(interface, "4.2.5 KneeAssist AI decision-support prototype (Objective 5)")
    reorder_blocks(document, model, [features, evaluation, data, interface, next_section])

    model_detail = find_paragraph(document, "The automated classification model was trained")
    set_text(
        model_detail,
        "The transfer-learning classification model was trained on prepared MRNet knee MRI studies and produces study-level scores "
        "for general abnormality, ACL tear and meniscal tear. ImageNet-pretrained ResNet-18 and EfficientNet-B0 encoders were "
        "implemented with a slice-aggregation head. The selected frozen EfficientNet-B0 checkpoint achieved the stronger recorded "
        "validation selection result. The objective of developing an automated study-level model was achieved, while the boundaries "
        "of its MRNet-only training data remain explicit."
    )
    feature_detail = find_paragraph(document, "Grad-CAM was used")
    set_text(
        feature_detail,
        "ResNet-18 and EfficientNet-B0 were compared using validation AUROC during model selection. EfficientNet-B0 was retained "
        "because it produced the stronger recorded validation result. Additional five-family benchmark candidates did not justify "
        "replacing the active checkpoint. This is a model-selection result from the available development data, not proof that one "
        "architecture will generalise best to every clinical scanner or patient population."
    )
    label_detail = find_paragraph(document, "The system was compared with available")
    set_text(
        label_detail,
        "MRI preparation began with 1,250 MRNet studies separated into 804 fitting, 226 tuning, 100 calibration and 120 official "
        "validation studies. Available image planes were normalised, resized, sampled consistently and aggregated at study level. "
        "Study-level and exact-series checks were implemented to reduce leakage and incompatible-input errors. Patient-level independence "
        "could not be established because patient linkage was unavailable."
    )
    preprocessing_detail = find_paragraph(document, "The external resources were useful")
    set_text(
        preprocessing_detail,
        "New inputs are inspected before prediction for supported arrays, NIfTI volumes, DICOM series and ZIP archives. The intake route "
        "checks file type, readable image content, slice count and available plane metadata. Where metadata cannot determine the plane, "
        "the interface requests confirmation instead of silently guessing. JPEG and PNG picture prediction are excluded because the model "
        "was not trained or validated on single internet images."
    )
    evaluation_detail = find_paragraph(document, "General abnormality is the primary result")
    set_text(
        evaluation_detail,
        "Performance was evaluated against the available MRNet official-validation reference labels, with KneeMRI ACL labels and a limited "
        "fastMRI meniscal subset retained only as separate external comparisons. General abnormality is the primary result for the narrowed "
        "academic topic and had the strongest observed ranking and threshold performance. Meniscal tear was weaker, with precision of 0.593 "
        "and sensitivity of 0.673; it remains an additional research output, not the main normal/abnormal claim. The macro accuracy of 0.764 "
        "must not be advertised as the probability that every finding for a new patient is correct. It averages three binary accuracies on this particular cohort."
    )
    interface_detail = find_paragraph(document, "The Streamlit workflow supports a case reference")
    set_text(
        interface_detail,
        "The Streamlit workflow supports a case reference, MRI array, NIfTI, DICOM or ZIP upload, intake inspection, manual plane "
        "confirmation where metadata is absent, analysis, slice viewing, prediction display, Grad-CAM attention, export and clearing. "
        "The supported ZIP route includes arrays, NIfTI volumes and guarded DICOM series. JPEG/PNG viewing and prediction are excluded. "
        "Browser verification confirmed the real-case upload and analysis path, attention display, exported JSON and clearing behaviour. "
        "Readable empty-input handling and dedicated NPY, NIfTI and DICOM input-safety tests were also verified. For every Grad-CAM view, "
        "the interface identifies the selected finding, MRI sequence and slice number, then describes the strongest heatmap concentration "
        "in displayed-image coordinates only. This helps the reviewer understand what score is being explained without claiming a named "
        "anatomical structure, confirmed tear or lesion boundary. This completes the decision-support prototype objective for supported research inputs."
    )

    set_text(
        find_paragraph(document, "The five objectives were addressed"),
        "The five objectives were addressed with different types of evidence. Controlled MRI preprocessing, transfer-learning classification, "
        "validation-based model selection, quantitative metrics against available labels and interface verification are observable technical "
        "achievements. The absence of a radiologist-reader study and weak external transfer are findings to report, not reasons to hide the results."
    )
    set_text(
        find_paragraph(document, "The objectives were deliberately specific and measurable"),
        "The objectives were deliberately specific and measurable. They are appropriate to a fourth-year computing project because each is "
        "supported by an implemented artifact or measured experiment. They do not promise replacement of specialists, a radiologist-reader "
        "study that was not available, or deployment certification. This allows the report to describe completed work without suppressing the observed weaknesses."
    )
    document.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
