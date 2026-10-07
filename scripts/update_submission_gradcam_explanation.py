"""Keep the submission report aligned with the dashboard Grad-CAM explanation."""

from pathlib import Path

from docx import Document


REPORT = Path(r"C:\Users\HP\Desktop\KneeAssist-AI\docs\academic\submission\Mabira_Mathew_R234371Y.docx")
ADDITION = (
    " For every Grad-CAM view, the interface identifies the selected finding, MRI sequence and slice number, "
    "then describes the strongest heatmap concentration in displayed-image coordinates only. This helps the reviewer "
    "understand what score is being explained without claiming a named anatomical structure, confirmed tear or lesion boundary."
)


def main() -> None:
    document = Document(REPORT)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("The Streamlit workflow supports a case reference"):
            if "strongest heatmap concentration" not in paragraph.text:
                paragraph.add_run(ADDITION)
            document.save(REPORT)
            print(REPORT)
            return
    raise RuntimeError("Dashboard verification paragraph was not found in Chapter 4.")


if __name__ == "__main__":
    main()
