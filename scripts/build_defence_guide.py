
# placeholder

from pathlib import Path
import json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r"C:\Users\HP\Desktop\KneeAssist-AI")
OUT=ROOT/"docs"/"academic"/"defence"
ASSETS=OUT/"KneeAssist_XAI_Model_and_Dashboard_Explainer_assets"
OUT.mkdir(parents=True,exist_ok=True);ASSETS.mkdir(parents=True,exist_ok=True)
REG=json.loads((ROOT/"model_registry"/"active_models.json").read_text(encoding="utf-8"))
FINAL=json.loads((ROOT/"runs"/"model_portfolio_v1"/"final"/"metrics.json").read_text(encoding="utf-8"))
SPLITS=json.loads((ROOT/"data"/"splits"/"SPLIT_MANIFEST.json").read_text(encoding="utf-8"))
NAVY="#173B59"; BLUE="#2E75B6"; PALE="#EAF2F8"; LIGHT="#F6F8FA"; GREY="#D9D9D9"; DARK="#222222"; ORANGE="#D96C00"; GREEN="#2F6B45"; WHITE="#FFFFFF"

def ft(size,bold=False):
    names=(["C:\\Windows\\Fonts\\arialbd.ttf","C:\\Windows\\Fonts\\calibrib.ttf"] if bold else [])+["C:\\Windows\\Fonts\\arial.ttf","C:\\Windows\\Fonts\\calibri.ttf"]
    for n in names:
        if Path(n).exists(): return ImageFont.truetype(n,size)
    return ImageFont.load_default()
def box(d,b,fill,outline=NAVY): d.rounded_rectangle(b,radius=22,fill=fill,outline=outline,width=3)
def wrap(d,xy,text,width,f,fill=DARK,space=6,center=False):
    words=text.split(); lines=[]; cur=""
    for w in words:
        test=w if not cur else cur+" "+w
        if d.textbbox((0,0),test,font=f)[2]<=width: cur=test
        else: lines.append(cur);cur=w
    if cur: lines.append(cur)
    y=xy[1]
    for line in lines:
        w=d.textbbox((0,0),line,font=f)[2];x=xy[0]+((width-w)/2 if center else 0)
        d.text((x,y),line,font=f,fill=fill);y+=d.textbbox((0,0),line,font=f)[3]+space
    return y
def arrow(d,a,b,color=BLUE):
    d.line([a,b],fill=color,width=6);ang=math.atan2(b[1]-a[1],b[0]-a[0]);s=18
    d.polygon([b,(b[0]-s*math.cos(ang-math.pi/6),b[1]-s*math.sin(ang-math.pi/6)),(b[0]-s*math.cos(ang+math.pi/6),b[1]-s*math.sin(ang+math.pi/6))],fill=color)
def image_flow():
    im=Image.new("RGB",(1800,850),"white");d=ImageDraw.Draw(im);d.text((48,30),"KneeAssist XAI live workflow",font=ft(42,True),fill=NAVY)
    data=[("1 MRI study","DICOM, NIfTI, NPY or a study ZIP"),("2 Intake checks","File type, dimensions, numeric values and plane evidence"),("3 Plane confirmation","Axial, coronal and sagittal; ask the user if uncertain"),("4 Model registry","Compatible candidates ranked using saved validation evidence"),("5 Study scores","DenseNet abnormal, ResNet ACL and Swin meniscus"),("6 XAI review","MRI slices, Grad-CAM attention and export")]
    xs=[35,330,625,920,1215,1510]
    for i,(h,b) in enumerate(data):
        x=xs[i];box(d,(x,170,x+250,490),PALE);wrap(d,(x+20,205),h,210,ft(27,True),NAVY,8);wrap(d,(x+20,300),b,210,ft(20),DARK,7)
        if i<5:arrow(d,(x+250,330),(xs[i+1]-12,330))
    d.text((48,590),"The uploaded study does not choose the winning model.",font=ft(30,True),fill=NAVY)
    wrap(d,(48,650),"After the MRI is inspected, the system retrieves the highest-ranked compatible registered route using saved validation evidence. It does not choose the largest patient-specific score and it does not retrain at upload time.",1660,ft(25),DARK,9)
    p=ASSETS/"workflow.png";im.save(p);return p
def image_planes():
    im=Image.new("RGB",(1800,820),"white");d=ImageDraw.Draw(im);d.text((48,30),"MRI planes used by the study-level model",font=ft(42,True),fill=NAVY)
    data=[("Axial","Cross-sectional layers through the knee. Think of looking from above or below.","Top and bottom slices"),("Coronal","Front-facing slice through the knee. Think of looking from the front.","Inner and outer sides"),("Sagittal","Side view through the knee. Think of looking from the side.","Structures along their length")]
    for i,(n,t,c) in enumerate(data):
        x=50+i*580;box(d,(x,140,x+510,735),LIGHT);d.text((x+165,170),n,font=ft(34,True),fill=NAVY);cx=x+255;cy=370
        if n=="Axial":
            d.ellipse((cx-145,cy-100,cx+145,cy+100),fill="#D6D6D6",outline=NAVY,width=4);d.line((cx-155,cy,cx+155,cy),fill=ORANGE,width=8)
            d.ellipse((cx-75,cy-48,cx-7,cy+20),fill="#FFFFFF",outline="#666666",width=3);d.ellipse((cx+7,cy-48,cx+75,cy+20),fill="#FFFFFF",outline="#666666",width=3)
        elif n=="Coronal":
            d.rounded_rectangle((cx-115,cy-135,cx+115,cy+135),radius=60,fill="#D6D6D6",outline=NAVY,width=4);d.line((cx,cy-150,cx,cy+150),fill=ORANGE,width=8)
        else:
            d.rounded_rectangle((cx-100,cy-135,cx+100,cy+135),radius=60,fill="#D6D6D6",outline=NAVY,width=4);d.line((cx-125,cy,cx+125,cy),fill=ORANGE,width=8)
        wrap(d,(x+40,560),c,430,ft(24,True),ORANGE,7,True);wrap(d,(x+40,625),t,430,ft(22),DARK,7,True)
    p=ASSETS/"planes.png";im.save(p);return p
def image_select():
    im=Image.new("RGB",(1800,960),"white");d=ImageDraw.Draw(im);d.text((48,28),"Validation-based target model selection",font=ft(42,True),fill=NAVY)
    for i,(n,t) in enumerate([("ResNet-18","Residual CNN"),("ResNet-50","Deeper residual CNN"),("DenseNet-121","Dense CNN"),("EfficientNet-B0","Efficient CNN"),("Swin-T","Window transformer")]):
        y=135+i*125;box(d,(55,y,410,y+85),PALE);d.text((78,y+14),n,font=ft(25,True),fill=NAVY);d.text((78,y+50),t,font=ft(18),fill=DARK)
    arrow(d,(435,430),(620,430));box(d,(630,240,1065,640),LIGHT);wrap(d,(675,285),"Same fixed MRNet protocol",345,ft(31,True),NAVY,10,True);wrap(d,(675,380),"Same study splits, target labels, 224 x 224 preprocessing, slice sampling and study-level aggregation.",345,ft(22),DARK,9,True);wrap(d,(675,525),"45 percent AUROC + 30 percent sensitivity + 15 percent average precision + 10 percent F1",345,ft(21,True),ORANGE,8,True);arrow(d,(1080,430),(1265,430))
    for i,(target,model,score) in enumerate([("General Abnormality","DenseNet-121","0.896"),("ACL Tear","ResNet-18","0.772"),("Meniscal Tear","Swin-T","0.795")]):
        y=135+i*205;box(d,(1280,y,1740,y+155),PALE);d.text((1310,y+20),target,font=ft(25,True),fill=NAVY);d.text((1310,y+63),model,font=ft(26,True),fill=GREEN);d.text((1310,y+110),"Saved validation score: "+score,font=ft(19),fill=DARK)
    wrap(d,(50,830),"Selection is performed before a new case is scored. The live system is target routing, not a probability ensemble.",1660,ft(24),DARK,8)
    p=ASSETS/"selection.png";im.save(p);return p
def image_cam():
    im=Image.new("RGB",(1800,750),"white");d=ImageDraw.Draw(im);d.text((48,28),"How to explain a Grad-CAM attention map",font=ft(42,True),fill=NAVY)
    for x,title,over in [(105,"Original MRI slice",False),(700,"Grad-CAM positive contribution overlay",True)]:
        box(d,(x,140,x+480,590),LIGHT)
        for yy in range(200,535):
            tone=int(35+90*math.exp(-((yy-365)/130)**2));d.line((x+50,yy,x+430,yy),fill=(tone,tone,tone))
        d.ellipse((x+125,240,x+355,500),fill=(115,115,115),outline=(220,220,220),width=3);d.ellipse((x+190,305,x+290,430),fill=(190,190,190),outline=(40,40,40),width=3)
        if over:
            for r,col in [(140,(215,75,0)),(95,(240,120,0)),(55,(255,190,15))]:d.ellipse((x+240-r,365-r,x+240+r,365+r),fill=col)
        d.text((x+40,620),title,font=ft(23,True),fill=NAVY)
    arrow(d,(610,370),(670,370),ORANGE);box(d,(1250,140,1735,590),PALE);wrap(d,(1292,195),"Correct interpretation",395,ft(31,True),NAVY,10,True);wrap(d,(1298,300),"Warm colours show image areas that contributed positively to the selected model score on this slice.",385,ft(24),DARK,9,True);wrap(d,(1298,455),"They do not confirm a tear, name a structure or define a lesion boundary.",385,ft(24,True),ORANGE,9,True)
    p=ASSETS/"gradcam.png";im.save(p);return p
FLOW=image_flow();PLANES=image_planes();SEL=image_select();CAM=image_cam()

def shade(cell,fill):
    p=cell._tc.get_or_add_tcPr();e=p.find(qn("w:shd"))
    if e is None:e=OxmlElement("w:shd");p.append(e)
    e.set(qn("w:fill"),fill.lstrip("#"))
def borders(cell,size="4"):
    p=cell._tc.get_or_add_tcPr();e=p.first_child_found_in("w:tcBorders")
    if e is None:e=OxmlElement("w:tcBorders");p.append(e)
    for n in ("top","bottom","left","right"):
        x=e.find(qn("w:"+n))
        if x is None:x=OxmlElement("w:"+n);e.append(x)
        x.set(qn("w:val"),"single");x.set(qn("w:sz"),size);x.set(qn("w:color"),GREY.lstrip("#"))
def margin(cell):
    p=cell._tc.get_or_add_tcPr();e=p.first_child_found_in("w:tcMar")
    if e is None:e=OxmlElement("w:tcMar");p.append(e)
    for n in ("top","start","bottom","end"):
        x=e.find(qn("w:"+n))
        if x is None:x=OxmlElement("w:"+n);e.append(x)
        x.set(qn("w:w"),"110");x.set(qn("w:type"),"dxa")
def table(doc,heads,rows,widths,sz=8.7):
    t=doc.add_table(rows=1,cols=len(heads));t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    tr=t.rows[0];tr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i,h in enumerate(heads):
        c=tr.cells[i];c.text=h;c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;shade(c,NAVY);margin(c);borders(c,"6")
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.bold=True;r.font.size=Pt(sz);r.font.color.rgb=RGBColor(255,255,255)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            c=cells[i];c.text=str(v);c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;shade(c,"F7FAFC" if ri%2==0 else WHITE);margin(c);borders(c)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(sz);r.font.color.rgb=RGBColor.from_string(DARK.lstrip('#'))
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
def head(doc,text,l=1):
    p=doc.add_heading(text,l);p._p.get_or_add_pPr().append(OxmlElement("w:keepNext"));return p
def para(doc,text,bold=None):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.12
    if bold and text.startswith(bold):r=p.add_run(bold);r.bold=True;p.add_run(text[len(bold):])
    else:p.add_run(text)
    for r in p.runs:r.font.size=Pt(10.6)
def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet");p.paragraph_format.space_after=Pt(3);p.add_run(text).font.size=Pt(10.4)
def number(doc,text):
    p=doc.add_paragraph(style="List Number");p.paragraph_format.space_after=Pt(4);p.add_run(text).font.size=Pt(10.4)
def cap(doc,text):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10);r=p.add_run(text);r.italic=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string("555555")

doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
for name,size in [("Normal",10.6),("Title",26),("Heading 1",17),("Heading 2",13),("Heading 3",11.5)]:
    s=doc.styles[name];s.font.name="Aptos";s._element.rPr.rFonts.set(qn("w:ascii"),"Aptos");s._element.rPr.rFonts.set(qn("w:hAnsi"),"Aptos");s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
doc.styles["Title"].font.bold=True
for n in ["Heading 1","Heading 2","Heading 3"]:doc.styles[n].font.bold=True
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=f.add_run("KneeAssist XAI Model and Dashboard Explainer | Research prototype | Clinical review required");r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string("666666")

p=doc.add_paragraph(style="Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(125);p.add_run("KneeAssist XAI Model and Dashboard Explainer")
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(20);r=p.add_run("A clear guide for explaining the MRI workflow, model selection, results and Grad-CAM");r.font.size=Pt(15)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(70);r=p.add_run("Prepared from the live application code, active model registry and saved MRNet evaluation evidence");r.font.size=Pt(11);r.italic=True
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(100);r=p.add_run("KneeAssist XAI V1");r.font.size=Pt(13);r.bold=True;r.font.color.rgb=RGBColor.from_string(NAVY.lstrip('#'))
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run("Research decision-support prototype. Not an autonomous diagnostic system.");r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string(ORANGE.lstrip('#'))
doc.add_page_break()

head(doc,"1. What this system is")
para(doc,"KneeAssist XAI is a local research decision-support prototype for study-level classification of knee MRI studies. It receives supported MRI volumes, checks their technical structure, identifies or asks the user to confirm the imaging planes, and produces three research scores: General Abnormality, ACL Tear and Meniscal Tear.")
para(doc,"The system is called XAI because it also produces Grad-CAM attention overlays. These overlays explain image regions that contributed positively to a selected model score. They are not a diagnosis, lesion boundary, detector or segmentation mask.")
head(doc,"What to say at the beginning of a presentation",2)
para(doc,"“My project is an explainable knee MRI research prototype. I compared five deep-learning architectures using one fixed MRNet protocol. The model registry stores the best validated compatible route for each target. When a new MRI study is uploaded, the system checks the input and uses those saved routes to produce study-level research scores and a Grad-CAM explanation.”")
head(doc,"Contents",2)
for x in ["1. What this system is","2. End-to-end workflow","3. MRI files, slices and planes","4. Preprocessing and study representation","5. Candidate algorithms and automatic selection","6. The live dashboard","7. Understanding results and warnings","8. Grad-CAM and explainable AI","9. Evaluation measures and recorded evidence","10. Plain-language glossary","11. Defence questions and answers","12. Limits and honest claims"]:bullet(doc,x)

doc.add_page_break();head(doc,"2. End-to-end workflow")
para(doc,"The live backend and frontend follow the sequence below. Training and model comparison happened offline during development. Uploading a new MRI does not start a training run.")
doc.add_picture(str(FLOW),width=Inches(6.85));cap(doc,"Figure 1. The implemented intake, model-routing and explainability workflow.")
head(doc,"How to demonstrate this in the application",2)
for x in ["Open Clinical MRI Review and upload one supported knee MRI study.","Open Study intake and plane confirmation. Explain that technical inspection happens before prediction.","Confirm any plane that lacks reliable DICOM or NIfTI orientation metadata.","Click Analyse study. Explain that the selector retrieves prior validation evidence rather than choosing the highest score from this patient.","Show the three finding cards, selected model information, MRI viewer and Grad-CAM tab.","Export the text or JSON research summary if a session record is needed."]:number(doc,x)

doc.add_page_break();head(doc,"3. MRI files, slices and planes")
para(doc,"An MRI study contains one or more sequences, and each sequence contains multiple slices. A slice is one two-dimensional cross-section. The model combines sampled slices across the available sequences to form one study-level representation.")
doc.add_picture(str(PLANES),width=Inches(6.85));cap(doc,"Figure 2. Orientation guide for the three MRI planes used by the application.")
head(doc,"Accepted input types",2)
table(doc,["Input","Meaning","How the application handles it"],[["DICOM (.dcm)","Clinical medical-imaging file with possible orientation metadata.","Reads geometry and series metadata, then groups slices into one series per plane."],["NIfTI (.nii/.nii.gz)","Three-dimensional medical-image volume.","Reads affine and voxel spacing; requests confirmation when orientation is uncertain."],["NumPy (.npy)","Processed numeric stack of MRI slices.","Requires a finite 3D array. A generic name such as 0002.npy needs manual plane confirmation."],["Study ZIP (.zip)","One package containing one supported MRI-study type.","Allows DICOM slices or up to three NIfTI/NPY volumes. Mixed formats are rejected."]],[1.28,2.05,3.40],8.7)
head(doc,"Technical checks",2)
for x in ["Each processed MRI sequence must be a finite numeric 3D array arranged as slices x height x width.","The supported range is 1 to 512 slices and image dimensions from 16 to 1024 pixels.","The application rejects zero-variation data, corrupted files, encrypted ZIP files, duplicate-plane stacks and mixed MRI studies.","JPEG and PNG pictures are rejected because the model was trained for MRI studies rather than internet pictures or screenshots.","File checks cannot prove that an uploaded study is a knee, belongs to one patient, or is clinically suitable. The user remains responsible for uploading the correct study."]:bullet(doc,x)

doc.add_page_break();head(doc,"4. Preprocessing and study representation")
para(doc,"The same core preprocessing approach was used for the saved research protocol and the live predictor. This consistency matters because evaluation would not be fair if uploaded data were processed differently from development data.")
table(doc,["Stage","What happens","Reason"],[["Slice sampling","Up to 12 evenly spaced slices are selected from each available plane.","Keeps a representative view while controlling computation."],["Intensity normalisation","The 1st and 99th intensity percentiles scale each sequence to a stable range.","Reduces raw intensity-range differences between sequences."],["Resize","Each sampled slice is resized to 224 x 224 pixels.","Matches the image size used by the encoders."],["Three-channel conversion","The grayscale slice is repeated into three channels and ImageNet-normalised.","Matches the input contract of ImageNet-pretrained encoders."],["Slice aggregation","The encoder produces features per slice. Mean and maximum pooling combine them within each plane.","Produces a study-level representation rather than a single-slice judgement."],["Plane mask","Missing planes are represented explicitly in the study head.","Allows incomplete studies with a visible caution."]],[1.35,3.1,2.28],8.4)
para(doc,"The standard presentation input contains axial, coronal and sagittal sequences. One-plane and two-plane inputs can be processed if technically compatible, but the dashboard warns that missing-sequence performance has not been separately validated.")
head(doc,"Important distinction",2);para(doc,"A single MRI slice may be shown in the viewer, but the three displayed scores are study-level outputs. They are calculated from sampled slices and available planes, not only the currently displayed image.")

doc.add_page_break();head(doc,"5. Candidate algorithms and automatic model selection")
para(doc,"Five architectures were compared using the same fixed MRNet protocol: ResNet-18, ResNet-50, DenseNet-121, EfficientNet-B0 and Swin Transformer. The aim was to compare them fairly and save the strongest validation route for each target.")
doc.add_picture(str(SEL),width=Inches(6.85));cap(doc,"Figure 3. The validation-evidence selection mechanism used by the live application.")
head(doc,"What each algorithm name means",2)
table(doc,["Model","Meaning","Role in this project"],[["ResNet-18","Residual convolutional neural network with 18 layers.","Candidate model; selected route for ACL Tear."],["ResNet-50","Deeper residual convolutional neural network with 50 layers.","Candidate model retained in the benchmark comparison."],["DenseNet-121","Dense convolutional network with 121 layers and feature reuse.","Candidate model; selected route for General Abnormality."],["EfficientNet-B0","Compact convolutional network scaled for efficient classification.","Candidate model retained in the benchmark comparison."],["Swin Transformer or Swin-T","Shifted Window Transformer, Tiny variant, using local-window attention.","Candidate model; selected route for Meniscal Tear."]],[1.65,3.0,2.08],8.5)
head(doc,"Selection score",2)
para(doc,"The registry uses a predefined weighted validation score: 45% AUROC, 30% sensitivity, 15% average precision and 10% F1-score. It is calculated from the fixed 226-study MRNet tuning partition. It does not use training accuracy, final official-validation results or the uploaded patient score.")
table(doc,["Target","Selected architecture","Saved validation score","Why it is used"],[["General Abnormality","DenseNet-121","0.896","Highest compatible saved weighted validation score."],["ACL Tear","ResNet-18","0.772","Highest compatible saved weighted validation score."],["Meniscal Tear","Swin Transformer","0.795","Highest compatible saved weighted validation score."]],[1.8,1.7,1.7,2.55],8.8)
para(doc,"At upload time, the system checks which benchmark candidates are compatible with the available planes. In the current registry, all five candidates support axial, coronal and sagittal planes and permit at least one plane. The registry returns the saved winner for each target. It never chooses a lower-ranked candidate silently if the winner is unavailable.")

doc.add_page_break();head(doc,"6. The live dashboard")
para(doc,"The Streamlit dashboard has three workspaces: Clinical MRI Review, Model Evaluation and Selection, and Research Workflow. The first analyses an upload. The second presents saved comparison evidence. The third gives a guided overview of the workflow.")
table(doc,["Dashboard area","What it shows","How to explain it"],[["Case reference","Optional text label for the research case.","It is only a display and export label. It is not model input; leaving it blank does not create a warning."],["MRI study files","Uploaded MRI sequences or study ZIP.","Use one study at a time. All files should belong to the same case."],["Study intake and plane confirmation","File type, shape, slice count, dimensions, plane evidence and required confirmation.","This is technical validity and compatibility checking, not diagnosis."],["Models used for this study","The selected architecture for each target.","These routes come from stored validation evidence, not patient scores."],["Finding cards","General Abnormality, ACL Tear and Meniscal Tear scores.","These are calibrated research scores, not definitive diagnoses."],["MRI slice viewer","A selected slice from a selected available plane.","It supports input review, not a full radiological reading."],["Model attention tab","A Grad-CAM overlay for one finding and one sequence.","It is coarse explanation of positive contribution to a score."],["Performance graphs","Saved ROC, precision-recall and confusion-matrix figures.","They show the comparison evidence before route registration."],["Export summary","Text and JSON download.","This exports displayed research output and routing details."]],[1.55,2.6,3.6],8.1)
head(doc,"Why the dashboard shows warnings",2)
for x in ["Input completeness caution means one or more planes is missing. It is not caused by an empty case reference.","Below research threshold means the score is below the stored cut-off. It does not rule out injury.","Clinical confidence not established means no independent clinical validation or radiologist reader study exists.","Detection boxes unavailable and segmentation masks unavailable are honest statements: no trained and validated localisation models are registered.","Attention could not be generated means only the optional visual explanation failed for that view. Study-level scores can remain available."]:bullet(doc,x)

doc.add_page_break();head(doc,"7. Understanding the three result cards")
table(doc,["Finding","What it refers to","What must not be claimed"],[["General Abnormality","The broad abnormal reference label in MRNet.","Do not say it identifies the precise injury, disease, treatment or affected tissue."],["ACL Tear","ACL means Anterior Cruciate Ligament, a stabilising ligament within the knee.","Do not say it confirms a tear or gives partial versus complete severity."],["Meniscal Tear","The meniscus is fibrocartilage that helps distribute force in the knee.","Do not say it identifies medial/lateral side, tear pattern or treatment."]],[1.65,3.1,2.95],8.8)
head(doc,"How a score becomes a displayed status",2)
para(doc,"Each selected route produces a deterministic calibrated score between 0 and 1. The application compares it with a stored route-specific threshold. General Abnormality uses 0.556, ACL Tear uses 0.129 and Meniscal Tear uses 0.310. The thresholds came from a reserved MRNet calibration partition.")
table(doc,["Displayed status","Meaning","Safe explanation"],[["Flagged for research review","Score is at or above its stored research threshold.","It supports research review; it is not a confirmed diagnosis."],["Below research threshold","Score is below its stored research threshold.","It was not flagged. This does not exclude injury."],["Near threshold","Score is close to its cut-off.","This is score distance only, not clinical confidence."],["MC-dropout variation","Variation from repeated exploratory classifier-head runs.","It is not injury severity, diagnosis or a measure of correctness."]],[1.65,2.95,3.1],8.5)
head(doc,"Example wording",2);para(doc,"“The ACL route used ResNet-18. Its calibrated study-level score was compared with the preselected ACL research threshold. The displayed status supports research review only; it does not confirm or exclude an ACL injury.”")

doc.add_page_break();head(doc,"8. Grad-CAM and explainable AI")
para(doc,"Grad-CAM means Gradient-weighted Class Activation Mapping. The system calculates gradients for the selected model score and turns the positive contribution of image regions into a heatmap over one sampled MRI slice.")
doc.add_picture(str(CAM),width=Inches(6.85));cap(doc,"Figure 4. Correct interpretation of the model attention overlay.")
head(doc,"Correct explanation",2);para(doc,"“Warmer colours show areas of this displayed slice that contributed positively to the selected model score. They help me inspect what the classifier used, but they do not prove that a tear is present in the highlighted area.”")
head(doc,"Why it can fail or look broad",2)
for x in ["Grad-CAM is calculated for the selected architecture and target, not from all routes combined.","It works with low-resolution feature maps enlarged to the 224 x 224 displayed slice. Broad heatmaps are expected.","It describes contribution to a score, not an anatomical contour. A central heatmap must never be labelled as a confirmed structure or lesion.","If attention calculation fails, predictions remain separate so that a visualisation error does not erase the study-level results."]:bullet(doc,x)

doc.add_page_break();head(doc,"9. Evaluation measures and recorded evidence")
para(doc,"Candidate selection used the fixed tuning validation partition. The following final development figures were calculated on 120 MRNet official-validation studies after selection. This cohort had previous development exposure and is not an independent external or clinical test.")
rows=[]
for k,label in [("abnormal","General Abnormality"),("acl","ACL Tear"),("meniscus","Meniscal Tear"),("macro","Macro average")]:
    m=FINAL["metrics"][k];rows.append([label,f"{m['auroc']:.3f}",f"{m['average_precision']:.3f}",f"{m['sensitivity']:.3f}",f"{m['specificity']:.3f}",f"{m['precision']:.3f}",f"{m['f1']:.3f}",f"{m['accuracy']:.3f}"])
table(doc,["Target","AUROC","AP","Sensitivity","Specificity","Precision","F1","Accuracy"],rows,[1.43,.67,.60,.83,.80,.72,.53,.65],7.9)
head(doc,"Metric definitions",2)
table(doc,["Metric","Meaning","How to explain it"],[["AUROC","Area Under the Receiver Operating Characteristic curve.","Measures ranking of positive and negative studies across thresholds. 1.0 is perfect; 0.5 is chance-level ranking."],["AP","Average Precision, a precision-recall curve summary.","Useful where reference labels are not evenly balanced."],["Sensitivity or recall","Proportion of labelled positive studies that were flagged.","Higher sensitivity reduces missed positive reference labels but may create more false positives."],["Specificity","Proportion of labelled negative studies that were not flagged.","Higher specificity reduces false positives but may miss more positives."],["Precision","Proportion of flagged studies that were labelled positive.","Shows agreement with reference labels in that evaluation cohort."],["F1-score","Harmonic balance of precision and sensitivity.","A single measure penalising poor precision or poor sensitivity."],["Accuracy","Proportion of all studies classified correctly at the chosen threshold.","It can mislead when labels are imbalanced, so it is not used alone."],["Calibration","How closely displayed scores correspond to observed proportions in a development cohort.","The current calibration used 100 reserved MRNet studies; it is not established externally or clinically."]],[1.15,2.52,3.14],8.0)
head(doc,"Fixed split evidence",2)
table(doc,["Partition","Studies","Role"],[["Fitting",SPLITS["roles"]["fit"]["studies"],"Model fitting."],["Tuning validation",SPLITS["roles"]["tune"]["studies"],"Candidate comparison and selection score only."],["Calibration",SPLITS["roles"]["calibration"]["studies"],"Calibration mapping and thresholds."],["Official validation",SPLITS["roles"]["official_valid"]["studies"],"Final recorded development results; previously used in development."]],[2.0,1.0,3.8],8.8)
para(doc,"The split audit records zero study overlap and zero exact-series overlap across partitions. The local MRNet copy does not provide patient linkage, so patient-level independence cannot be verified.")

head(doc,"10. Plain-language glossary")
gloss=[("AI","Artificial Intelligence. Here it means trained image-classification models."),("XAI","Explainable Artificial Intelligence. Here it refers to model explanation with Grad-CAM."),("MRI","Magnetic Resonance Imaging."),("ACL","Anterior Cruciate Ligament."),("MCL","Medial Collateral Ligament. Mentioned in knee medicine, but not an active model output."),("PCL","Posterior Cruciate Ligament. Mentioned in knee medicine, but not an active model output."),("Meniscus","Fibrocartilage that helps distribute knee load. Its plural is menisci."),("CNN","Convolutional Neural Network. ResNet, DenseNet and EfficientNet are CNN families."),("Transformer","A model family using attention. Swin-T is the Transformer candidate."),("Pretrained","Encoder began from ImageNet weights. It does not make the system a general internet-picture classifier."),("Model registry","Saved JSON record of candidate evidence, selected checkpoints, thresholds and compatibility."),("Checkpoint","Saved trained model weights used during inference."),("Inference","Running an already trained model on a new study. It is not training."),("Validation","Development data used to compare candidates and make selection decisions."),("Threshold","Stored score cut-off used to display flagged or below-threshold status."),("Grad-CAM","Visual explanation of positive contribution to a selected classifier score."),("Detection","Approximate boxes around labelled findings. Not enabled in this system."),("Segmentation","Pixel-level mask output. Not enabled in this system."),("MC-dropout","Repeated exploratory dropout-based score variation. It is not clinical uncertainty.")]
table(doc,["Term","Meaning"],gloss,[1.55,5.4],8.5)

doc.add_page_break();head(doc,"11. Defence questions and answers")
qas=[("What problem does the project solve?","It demonstrates a transparent workflow for reviewing a knee MRI study with validated image-classification routes and an attention map. It is research decision support, not autonomous diagnosis."),("Why use several algorithms?","Different architectures can perform differently for different targets. I compared five candidates fairly on the same saved protocol, then registered the strongest validation route for each target."),("Does the system choose a model after upload?","Yes. It checks input compatibility after upload and reads saved validation evidence from the registry. It does not choose the model with the biggest patient-specific score."),("Does it train while a doctor waits?","No. Training, comparison and calibration are offline development tasks. The live app only performs intake checks, routing, inference, attention generation and export."),("Why are three planes useful?","They give complementary views through the knee. The model combines sampled slices across axial, coronal and sagittal sequences into one study-level representation."),("What does a 70 percent result mean?","It is a calibrated development-model score, not a diagnosis or a probability that the model is correct. It must be read with its research threshold and clinical review."),("Why is the ACL threshold lower than 50 percent?","Thresholds were selected from reserved calibration data, not set simply at 50 percent. They balance sensitivity and specificity within the development protocol."),("What does Grad-CAM show?","It shows image regions that contributed positively to one selected model score on one displayed slice. It does not prove a lesion or identify a tear boundary."),("Can it locate or segment a tear?","No. There is no trained and validated detection-box or segmentation-mask model in the current registry."),("Can it recommend treatment?","No. Treatment requires clinical assessment and evidence beyond this research prototype."),("What is the main limitation?","Training used MRNet only. There is no independent clinical validation, patient linkage verification, reader study, detection model, segmentation model or treatment function.")]
for q,a in qas:
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(3);r=p.add_run("Question: "+q);r.bold=True;r.font.size=Pt(10.4)
    para(doc,"Answer: "+a,"Answer: ")

doc.add_page_break();head(doc,"12. Limits and honest claims")
head(doc,"What the system does",2)
for x in ["Accepts supported MRI studies in DICOM, NIfTI, NPY and compatible ZIP formats.","Checks technical file validity, dimensions, plane evidence and duplicate-plane conflicts.","Combines sampled slices across the available axial, coronal and sagittal planes.","Retrieves prevalidated target-specific model routes from the active registry.","Displays research scores, thresholds, routing information, MRI slices, Grad-CAM attention and exportable summaries."]:bullet(doc,x)
head(doc,"What the system does not do",2)
for x in ["Diagnose disease independently or replace a radiologist.","Confirm a tear, state its type, grade or severity, or rule out injury after a below-threshold score.","Classify random JPEG or PNG pictures, internet screenshots or photographs.","Identify MCL, PCL, fracture, effusion, osteoarthritis, cartilage damage or every knee condition.","Provide validated lesion boxes, segmentation masks, treatment recommendations or clinical confidence.","Claim independent external, clinical, prospective or radiologist-reader validation."]:bullet(doc,x)
head(doc,"Evidence checked for this guide",2)
table(doc,["Project artefact","What it verifies"],[["app.py","Visible dashboard workflow, supported input notice, score cards, model-used text, Grad-CAM, exports and warnings."],["src/data/preprocessing.py and smart_input_router.py","MRI formats, volume constraints, plane metadata logic, manual confirmation and preprocessing."],["src/models/study_model.py and aggregation.py","Slice feature extraction, mean and maximum aggregation and study-level head."],["src/selection/model_selector.py and routed_predictor.py","Compatibility filtering, saved validation-based routing and guard against patient-score selection."],["model_registry/active_models.json","Five-candidate population, active routes, thresholds, checkpoints and validation selection metrics."],["runs/model_portfolio_v1/final/metrics.json","Recorded selected-route final development metrics on 120 MRNet studies."]],[3.3,3.65],8.3)
para(doc,"Use this guide with the live Model Evaluation and Selection workspace. That page displays the candidate metric tables, rankings, ROC curves, precision-recall curves and confusion matrices that support the routing decision.")
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(40);r=p.add_run("End of guide");r.bold=True;r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string(NAVY.lstrip('#'))

doc.core_properties.title="KneeAssist XAI Model and Dashboard Explainer";doc.core_properties.subject="Defence guide for the current KneeAssist XAI research prototype";doc.core_properties.author="KneeAssist XAI Project"
target=OUT/"KneeAssist_XAI_Model_and_Dashboard_Explainer.docx";doc.save(target);print(target)





