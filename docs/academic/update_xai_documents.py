"""Create XAI-labelled, model-policy-aligned copies of the academic documents."""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SUBMISSION = ROOT / "docs" / "academic" / "submission"
REPORT_TITLE = (
    "KneeAssist XAI: Development of an Explainable Deep Learning System for "
    "Study-Level Classification of Normal and Abnormal Knee MRI Studies"
)


def paragraphs_in(document: Document):
    """Yield paragraphs from the main body, tables, headers, and footers."""
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def replace_brand(document: Document) -> None:
    for paragraph in paragraphs_in(document):
        for run in paragraph.runs:
            run.text = run.text.replace("KneeAssist AI", "KneeAssist XAI")
            # The launcher filename is deliberately retained for compatibility.
            run.text = run.text.replace("Launch KneeAssist XAI.cmd", "Launch KneeAssist AI.cmd")


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_starting(document: Document, start: str, replacement: str) -> bool:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(start):
            set_text(paragraph, replacement)
            return True
    return False


def remove_starting(document: Document, start: str) -> bool:
    """Remove a body paragraph that no longer belongs in the final report."""
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(start):
            paragraph._element.getparent().remove(paragraph._element)
            return True
    return False


def replace_starting_any(document: Document, start: str, replacement: str) -> bool:
    """Replace a paragraph beginning across body, tables, headers or footers."""
    for paragraph in paragraphs_in(document):
        if paragraph.text.startswith(start):
            set_text(paragraph, replacement)
            return True
    return False


def replace_starting_preserving_fields(document: Document, start: str, replacement: str) -> bool:
    """Replace caption/list wording without discarding a PAGEREF field."""
    for paragraph in paragraphs_in(document):
        if not paragraph.text.startswith(start):
            continue
        for run in paragraph.runs:
            if start in run.text:
                run.text = run.text.replace(start, replacement, 1)
                return True
        set_text(paragraph, replacement)
        return True
    return False


def replace_text_preserving_fields(document: Document, old: str, new: str) -> None:
    """Change visible caption text while retaining bookmarks and page fields."""
    for paragraph in paragraphs_in(document):
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def insert_paragraph_after(paragraph, text: str, style: str | None = None):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    result = Paragraph(element, paragraph._parent)
    if style:
        result.style = style
    result.add_run(text)
    return result


def add_bookmark(document: Document, paragraph, name: str) -> None:
    """Add a stable page-reference target for the static List of Tables."""
    ids = []
    for node in document._element.iter(qn("w:bookmarkStart")):
        value = node.get(qn("w:id"))
        if value and value.lstrip("-").isdigit():
            ids.append(int(value))
    bookmark_id = str(max(ids, default=0) + 1)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def append_page_reference(paragraph, bookmark_name: str) -> None:
    """Append an updateable PAGEREF field using Word OOXML primitives."""
    tab = paragraph.add_run()
    tab.add_tab()
    for kind, text in (("begin", None), ("instruction", f" PAGEREF {bookmark_name} \\h "), ("separate", None)):
        run = OxmlElement("w:r")
        if kind == "instruction":
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = text
            run.append(instruction)
        else:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), kind)
            run.append(field)
        paragraph._p.append(run)
    display = OxmlElement("w:r")
    display_text = OxmlElement("w:t")
    display_text.text = "0"
    display.append(display_text)
    paragraph._p.append(display)
    finish = OxmlElement("w:r")
    finish_field = OxmlElement("w:fldChar")
    finish_field.set(qn("w:fldCharType"), "end")
    finish.append(finish_field)
    paragraph._p.append(finish)


def insert_list_of_tables_entry(document: Document, text: str, bookmark_name: str) -> None:
    anchor = next(
        (
            paragraph for paragraph in document.paragraphs
            if paragraph.style.name == "toc 1" and paragraph.text.startswith("Table 4.2 ")
        ),
        None,
    )
    if anchor is None:
        raise RuntimeError("The List of Tables entry for Table 4.2 was not found")
    entry = insert_paragraph_after(anchor, "", anchor.style.name)
    entry.add_run(text)
    append_page_reference(entry, bookmark_name)


def append_field(paragraph, instruction: str, display: str) -> None:
    """Append a Word field with a visible fallback result.

    Word updates the displayed result during the final PDF export.  The fallback
    value keeps the document readable in renderers that do not calculate fields.
    """
    for kind, text in (("begin", None), ("instruction", instruction), ("separate", None)):
        run = OxmlElement("w:r")
        if kind == "instruction":
            instruction_node = OxmlElement("w:instrText")
            instruction_node.set(qn("xml:space"), "preserve")
            instruction_node.text = text
            run.append(instruction_node)
        else:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), kind)
            run.append(field)
        paragraph._p.append(run)
    visible = OxmlElement("w:r")
    visible_text = OxmlElement("w:t")
    visible_text.text = display
    visible.append(visible_text)
    paragraph._p.append(visible)
    finish = OxmlElement("w:r")
    finish_field = OxmlElement("w:fldChar")
    finish_field.set(qn("w:fldCharType"), "end")
    finish.append(finish_field)
    paragraph._p.append(finish)


def ensure_dynamic_page_numbers(document: Document, *, prefix: str = "") -> None:
    """Replace static footer numbers with real PAGE fields.

    The original departmental template carried literal ``i`` and ``1`` footer
    text.  That looks correct only on its first page, so the report receives a
    Roman PAGE field for preliminary matter and an Arabic PAGE field for the
    numbered report.  The short defence sheet uses a ``Page n`` footer.
    """
    for section_index, section in enumerate(document.sections):
        footer = section.footer
        current = " ".join(paragraph.text for paragraph in footer.paragraphs).strip()
        # The title section deliberately has no footer and is linked to no
        # independent footer part in the supplied template.
        if section_index == 0 and not current and footer.is_linked_to_previous:
            continue
        if footer.is_linked_to_previous:
            footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if prefix:
            paragraph.add_run(prefix)
        roman = current.lower() in {"i", "ii", "iii", "iv", "v"}
        append_field(paragraph, " PAGE \\* ROMAN " if roman else " PAGE ", "I" if roman else "1")
        for extra in footer.paragraphs[1:]:
            extra._element.getparent().remove(extra._element)


def remove_following_toc_entries(document: Document, heading_text: str) -> Paragraph:
    """Clear the existing static entries below a caption-list heading."""
    paragraphs = document.paragraphs
    heading_index = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == heading_text),
        None,
    )
    if heading_index is None:
        raise RuntimeError(f"Caption-list heading was not found: {heading_text}")
    heading = paragraphs[heading_index]
    for paragraph in paragraphs[heading_index + 1 :]:
        if paragraph.style.name != "toc 1":
            break
        paragraph._element.getparent().remove(paragraph._element)
    return heading


def rebuild_caption_list(document: Document, heading_text: str, captions, bookmark_prefix: str) -> None:
    """Build a sequential List of Figures or Tables with live page references."""
    heading = remove_following_toc_entries(document, heading_text)
    previous = heading
    for index, caption in enumerate(captions, start=1):
        bookmark_name = f"{bookmark_prefix}{index:02d}"
        add_bookmark(document, caption, bookmark_name)
        entry = insert_paragraph_after(previous, "", "toc 1")
        entry.add_run(caption.text)
        append_page_reference(entry, bookmark_name)
        previous = entry


def rebuild_caption_lists(document: Document) -> None:
    """Refresh both caption lists after routing-table and figure insertion."""
    figure_captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Figure Caption" and paragraph.text.startswith("Figure ")
    ]
    table_captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Table Caption" and paragraph.text.startswith("Table ")
    ]
    rebuild_caption_list(document, "LIST OF FIGURES", figure_captions, "KAFig")
    rebuild_caption_list(document, "LIST OF TABLES", table_captions, "KATbl")


def insert_project_figures(document: Document) -> None:
    """Insert the saved project figures beside their existing report captions."""
    figures = {
        "Figure 4.1": (
            ROOT / "runs/mri_finetune_02/results/resnet18/training_validation_curves.png",
            "Figure 4.1 Recorded ResNet-18 fine-tuning history. Source: project training artifacts. The curve describes this run and does not establish multi-seed stability.",
        ),
        "Figure 4.2": (
            ROOT / "runs/mri_finetune_02/results/efficientnet_b0/training_validation_curves.png",
            "Figure 4.2 Recorded EfficientNet-B0 fine-tuning history. Source: project training artifacts. Fine-tuning did not exceed the selected frozen warm-up checkpoint.",
        ),
        "Figure 4.3": (
            ROOT / "runs/mri_finetune_02/results/final/roc_curves.png",
            "Figure 4.3 ROC curves on the 120 MRNet official-validation studies. Source: project evaluation artifacts, 27 September 2026.",
        ),
        "Figure 4.4": (
            ROOT / "runs/mri_finetune_02/results/final/precision_recall_curves.png",
            "Figure 4.4 Precision–recall curves for the same MRNet cohort. Source: project evaluation artifacts.",
        ),
        "Figure 4.5": (
            ROOT / "runs/mri_finetune_02/results/final/confusion_matrices.png",
            "Figure 4.5 Confusion matrices at the formal baseline research thresholds. Source: project evaluation artifacts. Rows represent reference labels and columns predictions.",
        ),
        "Figure 4.6": (
            ROOT / "runs/mri_finetune_02/results/external_kneemri/roc_curves.png",
            "Figure 4.6 KneeMRI ACL external-reference ROC curve. Source: project evaluation artifacts. This previously evaluated cohort is not a new untouched test.",
        ),
        "Figure 4.7": (
            ROOT / "runs/mri_finetune_02/results/external_fastmri/roc_curves.png",
            "Figure 4.7 Exploratory fastMRI meniscal-reference ROC curve. Source: project evaluation artifacts. Annotation absence is not a definitive normal reference.",
        ),
        "Figure 4.8": (
            ROOT / "runs/mri_finetune_02/results/final/example_1130.png",
            "Figure 4.8 Saved MRI example and attention visualisation. Source: project inference artifact. The overlay is model attention, not a confirmed lesion.",
        ),
    }
    for paragraph in list(document.paragraphs):
        match = next((item for item in figures.items() if paragraph.text.startswith(item[0])), None)
        if match is None:
            continue
        _, (image_path, caption) = match
        if not image_path.is_file():
            raise FileNotFoundError(f"Verified report figure is missing: {image_path}")
        set_text(paragraph, caption)
        image_paragraph = paragraph.insert_paragraph_before()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.keep_with_next = True
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))


def renumber_following_chapter_four_tables(document: Document) -> None:
    """Reserve Table 4.3 for routing and retain consecutive Arabic numbering."""
    changes = (
        ("Table 4.3 Label distribution", "Table 4.4 Label distribution"),
        ("Table 4.4 External-reference results", "Table 4.5 External-reference results"),
        ("Table 4.5 Study-bootstrap AUROC intervals", "Table 4.6 Study-bootstrap AUROC intervals"),
        (
            "Table 4.6 Earlier baseline and selected release",
            "Table 4.7 Earlier baseline and formal EfficientNet-B0 baseline release",
        ),
    )
    for old, new in changes:
        replace_text_preserving_fields(document, old, new)


def format_routing_table(table) -> None:
    """Keep the seven-column evidence table readable at the guideline table size."""
    widths = (Inches(1.75), Inches(0.78), Inches(0.68), Inches(0.80), Inches(0.80), Inches(0.80), Inches(0.82))
    table.autofit = False
    for row in table.rows:
        for column, cell in enumerate(row.cells):
            cell.width = widths[column]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)


def insert_paragraph_after_table(table, text: str, style: str | None = None):
    """Insert a caption after a table while retaining the report's original order."""
    element = OxmlElement("w:p")
    table._tbl.addnext(element)
    paragraph = Paragraph(element, table._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_current_routing_table(document: Document) -> None:
    """Keep formal-baseline results distinct from current dashboard routes."""
    anchor = None
    for table in document.tables:
        values = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(values) >= 2 and values[0][:2] == ["Finding", "AUROC"] and values[1][0] == "General abnormality" and values[1][1] == "0.901":
            anchor = table
            break
    if anchor is None:
        raise RuntimeError("The formal baseline performance table was not found")
    if not replace_starting_preserving_fields(
        document,
        "Table 4.2 MRNet official validation performance",
        "Table 4.2 Formal EfficientNet-B0 baseline performance on the MRNet official-validation partition",
    ):
        raise RuntimeError("The formal baseline caption was not found")
    if not replace_starting(
        document,
        "The following evaluation paragraph returns",
        "Table 4.2 is the formal EfficientNet-B0 baseline evaluation used for the required ResNet-18-versus-EfficientNet-B0 comparison objective. It is not the performance table for the later dashboard routing registry. Its performance was evaluated against the available MRNet official-validation reference labels, with KneeMRI ACL labels and a limited fastMRI meniscal subset retained only as separate historical external comparisons. General abnormality is the primary academic result; meniscal scoring remains an additional research output. The baseline macro accuracy must not be interpreted as the probability that every finding for a new patient is correct.",
    ):
        raise RuntimeError("The formal baseline interpretation paragraph was not found")
    renumber_following_chapter_four_tables(document)
    caption = insert_paragraph_after_table(
        anchor,
        "Table 4.3 Current fixed target-routing development metrics on the previously used MRNet official-validation partition",
        "Table Caption",
    )
    current = document.add_table(rows=1, cols=7)
    current.style = anchor.style
    headers = ["Finding / active route", "AUROC", "F1", "Sensitivity", "Specificity", "Precision", "Accuracy"]
    for cell, value in zip(current.rows[0].cells, headers):
        cell.text = value
    rows = [
        ["General abnormality / DenseNet-121", "0.861", "0.929", "0.968", "0.560", "0.893", "0.883"],
        ["ACL tear / ResNet-18", "0.870", "0.763", "0.926", "0.591", "0.649", "0.742"],
        ["Meniscal tear / Swin-T", "0.808", "0.677", "0.846", "0.500", "0.564", "0.650"],
        ["Macro average", "0.846", "0.790", "0.914", "0.550", "0.702", "0.758"],
    ]
    for values in rows:
        cells = current.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    format_routing_table(current)
    caption._p.addnext(current._tbl)
    explanation = OxmlElement("w:p")
    current._tbl.addnext(explanation)
    paragraph = Paragraph(explanation, current._parent)
    paragraph.add_run("The dashboard selects these three routes from saved development evidence before an uploaded study is scored. It is not a patient-specific winner-takes-all selection and it is not a weighted ensemble. These values were measured on the same previously used 120-study MRNet development-validation cohort, so they are development evidence rather than independent or clinical generalisation.")


def update_report(source: Path, destination: Path) -> None:
    doc = Document(source)
    replace_brand(doc)
    set_text(doc.paragraphs[0], REPORT_TITLE)
    replacements = {
        "The implemented system is technically operational, but": (
            "The implemented system is technically operational, but its measured performance does not justify routine "
            "clinical deployment. The formal historical EfficientNet-B0 baseline achieved macro AUROC 0.8267 on the "
            "previously used 120-study MRNet official-validation cohort. The current dashboard uses separately fixed "
            "target routes and recorded macro AUROC 0.846 on that same development-validation cohort. Historical "
            "external-reference AUROCs of 0.6030 for ACL injury and 0.6561 for meniscal annotations belong only to "
            "the formal EfficientNet-B0 baseline. Neither evidence set is independent clinical validation. These "
            "results support further investigation rather than a claim of reliable diagnosis."
        ),
        "Knee magnetic resonance imaging classification requires more": (
            "Knee magnetic resonance imaging classification requires more than plausible model scores: "
            "data separation, reference labels and evaluation conditions determine what the scores can support. "
            "KneeAssist XAI is a local research application for study-level normal-versus-abnormal knee MRI "
            "classification. ACL and meniscal scores are secondary research outputs. MRNet was the sole training "
            "source: 804 fitting, 226 tuning, 100 calibration and 120 official-validation studies. The formal "
            "fourth-year comparison evaluated ImageNet-pretrained ResNet-18 and EfficientNet-B0 using validation "
            "AUROC; the EfficientNet-B0 baseline achieved 0.901 AUROC for general abnormality on the recorded MRNet "
            "official-validation cohort. A later portfolio benchmark supplied the current dashboard's fixed, "
            "prevalidated target routing: DenseNet-121 for general abnormality, ResNet-18 for ACL scoring and Swin-T "
            "for meniscal scoring. This is not patient-specific model selection or a weighted ensemble. The interface "
            "checks supported MRI input, displays sampled slices, model identity, scores and Grad-CAM attention, and "
            "can export a research summary. Attention is a coarse explanation of score contribution, not confirmed "
            "lesion localisation. Evidence remains MRNet-only, patient linkage was unavailable and external-reference "
            "evidence was limited; no radiologist-reader comparison was conducted. The result is an auditable research "
            "prototype, not a clinically validated diagnostic system."
        ),
        "The primary experimental comparison is narrower": (
            "The formal primary comparison is narrower than a benchmark competition. It compares ResNet-18 and "
            "EfficientNet-B0 under the project configuration and selects the baseline checkpoint by internal AUROC. "
            "A later supplementary frozen-encoder benchmark evaluated ResNet-50, DenseNet-121 and Swin Transformer "
            "under one locked split. The current dashboard uses the benchmark only as fixed target routing evidence: "
            "DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for meniscus when the input is "
            "compatible. It does not choose a model from the uploaded study's prediction scores and it does not use "
            "a weighted ensemble. There is no multi-seed cross-validation, prospective trial or exhaustive "
            "hyperparameter search."
        ),
        "CRISP-DM is used to structure": (
            "CRISP-DM structures six connected stages. Problem understanding defines the three findings and "
            "research-only use. Data understanding identifies supported labels. Preparation converts supported studies "
            "into reproducible inputs. Modelling performs the formal two-encoder comparison and preserves the later "
            "five-family portfolio separately. Evaluation examines internal and external-reference behaviour. Deployment "
            "packages a prevalidated target-routing registry in the local interface; the registry decision is made "
            "before a case is uploaded."
        ),
        "Both ResNet-18 and EfficientNet-B0 use": (
            "The formal baseline comparison uses ImageNet-pretrained ResNet-18 and EfficientNet-B0 slice encoders. "
            "Each study representation is formed by slice aggregation and presented to a multi-label head. The "
            "current dashboard registry holds separately trained, compatible models for its three displayed targets: "
            "DenseNet-121 (general abnormality), ResNet-18 (ACL) and Swin-T (meniscus). It selects each model from "
            "pre-upload validation evidence and available MRI planes, never from which model gives the highest score "
            "for the current case."
        ),
        "The transfer-learning classification model was trained": (
            "The study-level classification workflow was trained on prepared MRNet knee MRI studies and produces "
            "research scores for general abnormality, ACL tear and meniscal tear. The fourth-year baseline implemented "
            "ImageNet-pretrained ResNet-18 and EfficientNet-B0 with slice aggregation; EfficientNet-B0 achieved the "
            "stronger recorded baseline validation result. The current dashboard then applies a separate fixed routing "
            "registry from the portfolio evidence: DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T "
            "for meniscus. This extends the interface without changing the formal two-model comparison objective."
        ),
        "An example case used for workflow testing returned approximately": (
            "An earlier workflow test using the historical single-model baseline recorded one set of illustrative "
            "scores. Those values are retained only as an audit record and do not describe the current target-routed "
            "dashboard, which uses separate registered models for the three findings. Dashboard scores may vary with "
            "the selected registered route and exploratory MC-dropout sampling. No single demonstration case is a "
            "performance estimate or a diagnosis; the reported evaluation tables remain the appropriate evidence."
        ),
        "Performance was evaluated against the available MRNet official-validation reference labels": (
            "The following evaluation paragraph returns to the formal EfficientNet-B0 baseline. Its performance was "
            "evaluated against the available MRNet official-validation reference labels, with KneeMRI ACL labels and "
            "a limited fastMRI meniscal subset retained only as separate external comparisons. General abnormality is "
            "the primary result for the narrowed academic topic and had the strongest observed ranking and threshold "
            "performance. Meniscal tear was weaker, with precision of 0.593 and sensitivity of 0.673; it remains an "
            "additional research output, not the main normal/abnormal claim. The macro accuracy of 0.764 must not be "
            "advertised as the probability that every finding for a new patient is correct. It averages three binary "
            "accuracies on this particular cohort."
        ),
        "After the primary selection, a supplementary frozen-encoder benchmark": (
            "After the primary selection, a supplementary frozen-encoder benchmark compared five MRNet-only candidates "
            "on one locked internal-tuning split. EfficientNet-B0 was the strongest new macro candidate, but it did not "
            "exceed the historical calibrated EfficientNet-B0 baseline on that macro criterion, so no single overall "
            "replacement was claimed. Separately, the current dashboard registry records fixed target-level routes from "
            "saved development evidence: DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for meniscus. "
            "This routing is not a weighted ensemble and is never selected from an individual uploaded study's scores. "
            "These results are development evidence, not proof of clinical superiority."
        ),
        "The comparison includes a frozen-encoder warm-up": (
            "The formal ResNet-18-versus-EfficientNet-B0 comparison includes a frozen-encoder warm-up and attempted "
            "partial fine-tuning. Its historical baseline release is the EfficientNet-B0 frozen warm-up checkpoint "
            "followed by calibration. Describing it as a successful improved fine-tuned encoder would be inaccurate. "
            "The selection criterion was internal macro AUROC, not training accuracy or external performance. The "
            "separate live dashboard policy uses fixed target routes rather than this one baseline checkpoint for every score."
        ),
        "Results are organised by the five objectives": (
            "Results are organised by the five objectives and drawn from the preserved run artifacts. Figures reproduce "
            "project-generated outputs rather than illustrative performance. Numerical values are rounded for presentation. "
            "The formal historical baseline is the calibrated EfficientNet-B0 frozen warm-up checkpoint evaluated on the "
            "120-study MRNet official-validation partition. The live dashboard is a separate fixed target-routing policy: "
            "DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for meniscus. Neither evidence set is an "
            "independent clinical test."
        ),
        "The project demonstrates that an end-to-end knee MRI research application": (
            "The project demonstrates that an end-to-end knee MRI research application can be assembled around pretrained "
            "image encoders, study-level pooling and a modular inference pipeline. It also demonstrates that a higher "
            "internal selection score does not necessarily translate to higher performance on other evaluation material. "
            "The formal historical EfficientNet-B0 frozen warm-up release achieved MRNet macro AUROC 0.8267, while its "
            "external ACL and meniscal-reference AUROCs were 0.6030 and 0.6561. Separately, the live dashboard's fixed "
            "target-routing development metrics have macro AUROC 0.846 on the same previously used 120-study MRNet cohort. "
            "Neither is independent clinical validation."
        ),
        "ResNet-18 and EfficientNet-B0 were compared": (
            "ResNet-18 and EfficientNet-B0 were compared using validation AUROC for the formal model-comparison "
            "objective, and EfficientNet-B0 was retained as the stronger recorded baseline. The later five-family "
            "portfolio is used in the live dashboard only for target-specific, prevalidated routing: DenseNet-121 for "
            "general abnormality, ResNet-18 for ACL and Swin-T for meniscus. This routing is fixed before upload, is "
            "not a probability ensemble and never promotes a model because it scored highest for one patient. Neither "
            "baseline selection nor routing proves generalisation to every scanner or patient population."
        ),
        "Selected checkpoint: runs/mri_finetune_02/models/best_model.pth.": (
            "Formal baseline checkpoint: runs/mri_finetune_02/models/best_model.pth (EfficientNet-B0, MRNet-only, "
            "frozen warm-up followed by calibration). Current dashboard registry: model_registry/active_models.json "
            "routes DenseNet-121 to general abnormality, ResNet-18 to ACL and Swin-T to meniscus from pre-upload "
            "validation evidence. The registry is not an ensemble and does not use patient-specific scores to choose a model."
        ),
        "5. Retain the separate five-family benchmark": (
            "5. Retain the five-family portfolio as documented target-routing evidence. Update the registry only after "
            "a candidate meets the locked validation policy; never select a model because it gives the highest score "
            "for an uploaded study."
        ),
    }
    for start, replacement in replacements.items():
        if not replace_starting(doc, start, replacement):
            raise RuntimeError(f"Expected report paragraph was not found: {start[:60]}")
    insert_current_routing_table(doc)
    table_replacements = {
        "Developed a study-level EfficientNet-B0 transfer-learning classifier": (
            "Developed a study-level transfer-learning classifier with fixed target-routed dashboard models"
        ),
        "The final release evidence records 33 automated tests passing": (
            "The final release evidence records 43 automated tests passing in the final regression check and nine "
            "notebooks executed successfully"
        ),
        "Checkpoint SHA-256: b60293f6c7130fd646c81b6d758584faf3e17b2689503b5278ad711156d546da": (
            "Formal baseline checkpoint SHA-256: b60293f6c7130fd646c81b6d758584faf3e17b2689503b5278ad711156d546da. "
            "The live routed checkpoint identities and hashes are recorded in model_registry/active_models.json and "
            "models/checkpoint_manifest.json; all remain MRNet-only research artifacts."
        ),
        "EfficientNet-B0 retained as the selected model": (
            "EfficientNet-B0 retained as the formal historical baseline; live routes are reported separately"
        ),
        "6. Save and select the eligible checkpoint using internal macro AUROC only.": (
            "6. Select the formal baseline checkpoint using internal macro AUROC; maintain live target routes only through the locked registry policy."
        ),
        "3. Apply the same preprocessing and load the identified checkpoint.": (
            "3. Apply the same preprocessing and load the registered target-specific checkpoint for each reported finding."
        ),
    }
    for start, replacement in table_replacements.items():
        if not replace_starting_any(doc, start, replacement):
            raise RuntimeError(f"Expected report table or appendix text was not found: {start[:60]}")
    if not remove_starting(doc, "The bibliography contains"):
        raise RuntimeError("The bibliography explanatory paragraph was not found")
    insert_project_figures(doc)
    rebuild_caption_lists(doc)
    ensure_dynamic_page_numbers(doc)
    doc.core_properties.title = REPORT_TITLE
    doc.core_properties.subject = "Explainable deep learning for knee MRI study classification"
    doc.core_properties.keywords = "KneeAssist XAI; explainable AI; MRNet; knee MRI; research prototype"
    doc.save(destination)


def update_defence_sheet(source: Path, destination: Path) -> None:
    doc = Document(source)
    replace_brand(doc)
    set_text(doc.paragraphs[0], "KneeAssist XAI Defence Cheat Sheet")
    replacements = {
        "My project is KneeAssist XAI": (
            "My project is KneeAssist XAI, a local research prototype for study-level knee MRI classification. "
            "My formal academic model comparison used ResNet-18 and EfficientNet-B0; EfficientNet-B0 was the stronger "
            "recorded baseline. The current dashboard uses fixed target-specific routing from separately recorded "
            "validation evidence: DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for meniscus. "
            "It chooses these before the case is scored, so it is not selecting the model with the highest patient "
            "score and it is not an ensemble. Grad-CAM is coarse attention explaining score contribution, not confirmed "
            "anatomical localisation. The system is not a diagnostic, detection or segmentation system."
        ),
        "Answer: A separate five-family benchmark was completed": (
            "Answer: Yes. The portfolio compared five candidate architectures. The dashboard now uses only its "
            "prevalidated target entries: DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for "
            "meniscus. The selected entry is fixed from development evidence before upload. No weighted ensemble is "
            "active, and no model is chosen from a patient's prediction scores."
        ),
        "Click Analyse study: explain preprocessing, sampled slices, EfficientNet-B0": (
            "Click Analyse study: explain preprocessing, sampled slices, the target-specific prevalidated model and "
            "study-level aggregation."
        ),
    }
    for start, replacement in replacements.items():
        if not replace_starting(doc, start, replacement):
            raise RuntimeError(f"Expected defence paragraph was not found: {start[:60]}")
    defence_replacements = {
        "Use these exact numbers only with their dataset and limitation.": (
            "Use these exact values only as the formal historical EfficientNet-B0 baseline metrics on the previously "
            "used 120-study MRNet official-validation partition. The live dashboard has separate fixed target routes; "
            "its recorded routed development macro AUROC is 0.846 on the same previously used cohort. Neither result "
            "is independent clinical validation."
        ),
        "Question: Why EfficientNet-B0?": "Question: Why was EfficientNet-B0 retained as the formal baseline?",
        "Answer: It gave the stronger recorded internal-selection macro AUROC": (
            "Answer: In the formal ResNet-18-versus-EfficientNet-B0 comparison, it gave the stronger recorded "
            "internal-selection macro AUROC while remaining computationally manageable for a fourth-year prototype. "
            "That historical baseline is distinct from the live dashboard routes."
        ),
        "Answer: It checks the registry for a compatible checkpoint": (
            "Answer: It checks the registry for compatible prevalidated routes and shows their evidence. Before scoring, "
            "the dashboard assigns DenseNet-121 to general abnormality, ResNet-18 to ACL and Swin-T to meniscus. It "
            "does not choose a model from a patient's prediction scores and it is not a weighted ensemble."
        ),
        "An ImageNet-pretrained EfficientNet-B0 encoder converts every sampled MRI slice into image features.": (
            "The formal baseline and the live target routes use ImageNet-pretrained encoders. The live dashboard uses "
            "DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for meniscus."
        ),
        "ResNet-18 and EfficientNet-B0 were compared using internal-tuning macro AUROC.": (
            "ResNet-18 and EfficientNet-B0 were compared using internal-tuning macro AUROC; EfficientNet-B0 is the "
            "formal historical baseline. The later portfolio supplies fixed live routes: DenseNet-121 for general "
            "abnormality, ResNet-18 for ACL and Swin-T for meniscus."
        ),
        "Shows active checkpoint evidence and supported file types.": (
            "Shows active target-routing evidence, historical baseline provenance and supported file types."
        ),
        "Loads the compatible registered checkpoint, preprocesses the MRI and calculates study-level scores.": (
            "Loads the three compatible registered target models, preprocesses the MRI and calculates study-level scores."
        ),
        "The registry currently has one compatible validated research checkpoint": (
            "The registry has three fixed validated research routes. Each finding is assigned before upload scoring; this "
            "is safe target routing, not a fake multi-model ensemble."
        ),
        "EfficientNet produces a low-resolution final feature map": (
            "The selected classifier produces a low-resolution final feature map, which is enlarged to the 224 x 224 "
            "display. The model also has study-level labels, not lesion locations."
        ),
    }
    for start, replacement in defence_replacements.items():
        if not replace_starting_any(doc, start, replacement):
            raise RuntimeError(f"Expected defence table or paragraph was not found: {start[:60]}")
    if not replace_starting_any(
        doc,
        "Open model attention:",
        "Open model attention: choose a reported finding, then click Generate attention map. Explain that the overlay is "
        "coarse model attention for that route and slice, not a confirmed lesion location.",
    ):
        raise RuntimeError("Expected defence demonstration wording was not found")
    ensure_dynamic_page_numbers(doc, prefix="Page ")
    doc.core_properties.title = "KneeAssist XAI Defence Cheat Sheet"
    doc.core_properties.subject = "KneeAssist XAI model and dashboard defence notes"
    doc.save(destination)


def main() -> None:
    report_source = SUBMISSION / "Mabira_Mathew_R234371Y.docx"
    defence_source = SUBMISSION / "KneeAssist_AI_Defence_Cheat_Sheet.docx"
    update_report(report_source, SUBMISSION / "Mabira_Mathew_R234371Y_XAI.docx")
    update_defence_sheet(defence_source, SUBMISSION / "KneeAssist_XAI_Defence_Cheat_Sheet.docx")


if __name__ == "__main__":
    main()
