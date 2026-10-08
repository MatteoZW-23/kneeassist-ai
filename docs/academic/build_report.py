from pathlib import Path
import json, re, sys, unicodedata
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
doc = Document()
for st in doc.styles:
    if hasattr(st,'font'):
        st.font.name='Times New Roman'
for name in ['TOC 1','TOC 2','TOC 3']:
    if name not in doc.styles: doc.styles.add_style(name,1)
    st=doc.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(12); st.font.bold=False
    st.paragraph_format.line_spacing=1.15; st.paragraph_format.space_after=Pt(3)
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2.54)
    s.header_distance=s.footer_distance=Cm(1.27)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Caption']:
    st=doc.styles[name]
    st.font.name='Times New Roman'
    st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(12)
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.widow_control=True
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    st.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
for name in ['Heading 1','Heading 2','Heading 3','Heading 4']:
    st=doc.styles[name]; st.font.bold=True
    st.font.size=Pt(14 if name in ['Heading 1','Heading 2'] else 12)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before=Pt(12)
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Title'].font.size=Pt(18)
doc.styles['Title'].font.bold=True
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.line_spacing=1.15
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.core_properties.title='KneeAssist XAI development of an explainable deep learning system for study level classification of normal and abnormal knee MRI studies'
doc.core_properties.subject='Five objective academic project report'
doc.core_properties.author='Student details to be completed'
doc.core_properties.keywords='KneeAssist XAI; explainable AI; MRNet; research prototype; five objectives'

def p(text='',style=None): return doc.add_paragraph(text,style)
def field(par,code):
    r=par.add_run(); e=OxmlElement('w:fldSimple'); e.set(qn('w:instr'),code); r._r.addnext(e)
def number_section(sec,fmt,start):
    sec.footer.is_linked_to_previous=False
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(fp,'PAGE')
    num=OxmlElement('w:pgNumType'); num.set(qn('w:fmt'),fmt); num.set(qn('w:start'),str(start)); sec._sectPr.append(num)
def front_heading(text,newpage=True):
    if newpage: doc.add_page_break()
    q=p(text); q.style='Heading 1'; q.paragraph_format.page_break_before=False
    return q
def table(rows):
    t=doc.add_table(rows=1,cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    n=len(rows[0]); total=6.26
    widths=([1.5]+[(total-1.5)/(n-1)]*(n-1)) if n>=5 else ([1.38,2.42,2.46] if n==3 else [total/n]*n)
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(c,text) in enumerate(zip(cells,row)):
            c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            q=c.paragraphs[0]; q.add_run(text)
            q.paragraph_format.line_spacing=1.15; q.paragraph_format.space_before=Pt(4); q.paragraph_format.space_after=Pt(4)
            q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or n<=4 else WD_ALIGN_PARAGRAPH.CENTER
            for r in q.runs: r.font.size=Pt(10); r.bold=(i==0)
            pr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'80'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            if i==0:
                fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'E8E8E8'); pr.append(fill)
        trpr=t.rows[i]._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit'); trpr.append(ns)
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    p().paragraph_format.space_after=Pt(2)
    return t

# Cover and front matter
for text,style in [('MIDLANDS STATE UNIVERSITY',None),('Faculty of Business Sciences',None),('Department of Information and Marketing Sciences',None),('KneeAssist XAI development of an explainable deep learning system for study level classification of normal and abnormal knee MRI studies','Title'),('Academic project report',None),('Bachelor of Commerce Honours Degree in Data Science and Informatics',None),('Student name: [To be completed]',None),('Registration number: [To be completed]',None),('Supervisor: [To be completed]',None),('Project level: [2.2 or 4.2 — confirm]',None),('Campus: [Gweru or Harare — confirm]',None),('September 2026',None)]:
    q=p(text,style); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after=Pt(15 if style=='Title' else 10)
p('Prepared for student and supervisor review. Identity details, signatures and institutional submission requirements remain to be completed.').alignment=WD_ALIGN_PARAGRAPH.CENTER
sec=doc.add_section(WD_SECTION_START.NEW_PAGE); number_section(sec,'lowerRoman',1)
front_heading('Declaration',False)
p('This declaration is provided for the student to review and sign after checking the complete report, source material and project evidence. No signature or claim of unaided authorship has been supplied.')
p('I, [full name and registration number], confirm, upon signing, that I have reviewed this report and can account for its methods, results and cited sources. Assistance from generative AI is disclosed in Appendix A. I will not represent AI-assisted writing, coding or analysis as unaided work. I confirm that any additional assistance or previously submitted material will be declared before submission.')
p('Student signature: ____________________    Date: ____________________')
front_heading('Submission approval')
p('For completion by the authorised supervisor. This page is not evidence that approval has been granted.')
p('Project title: KneeAssist XAI development of an explainable deep learning system for study level classification of normal and abnormal knee MRI studies')
p('Student: [To be completed]    Registration number: [To be completed]')
p('Supervisor name: ____________________')
p('Approval decision and comments: __________________________________________')
p('Supervisor signature: ____________________    Date: ____________________')
front_heading('Abstract')
abstract=('Knee magnetic resonance imaging classification requires more than a model that produces plausible scores: data separation, reference labels and evaluation conditions determine what those scores can support. This project developed KneeAssist XAI, a local explainable-AI research application whose primary task is study-level classification of normal and abnormal knee MRI studies. ACL tear and meniscal tear scores are retained as secondary research outputs. Five objectives addressed prototype delivery, study preparation, transfer-learning modelling, formal comparison and performance evaluation. A retrospective computational design used MRNet as the sole training source, with 804 fitting, 226 internal tuning, 100 calibration and 120 official validation studies. ImageNet-pretrained ResNet-18 and EfficientNet-B0 encoders processed sampled slices from available imaging planes, followed by study-level aggregation and a multi-label classification head. The formal historical baseline was the frozen-encoder EfficientNet-B0 warm-up model; attempted fine-tuning did not exceed its internal selection performance. A later five-family portfolio supplies fixed prevalidated dashboard routes: DenseNet-121 for general abnormality, ResNet-18 for ACL and Swin-T for meniscus. The primary normal/abnormal output achieved AUROC of 0.901 on the official MRNet validation cohort. The interface supports prepared MRI arrays, cautious NIfTI input and guarded DICOM routes. Independent inference, Grad-CAM attention, MRI upload, export controls and case clearing were verified; the final regression check recorded 43 passing tests and nine executed notebooks. A separate RSNA Kaggle pilot was not merged because only 58 studies had complete explicit labels. Patient linkage was unavailable, evaluation material had been reused, and external reference evidence was limited. The contribution is an auditable working research prototype with transparent limitations and reproducible artifacts, rather than a clinically validated diagnostic system. Independent patient-level evaluation and controlled robustness studies are recommended before stronger claims.')
assert 250<=len(abstract.split())<=300,len(abstract.split())
p(abstract)
front_heading('Acknowledgements')
p('The project uses research resources and methods made available by the authors and custodians cited in the report. Personal acknowledgements should be added by the student after confirming the individuals and contributions involved. No supervisor contribution, funding award or institutional partnership is assumed.')
front_heading('Dedication')
p('[Optional personal dedication to be completed by the student.]')
front_heading('Table of contents')
field(p(),r'TOC \o "1-3" \h \z \u')
front_heading('List of figures')
field(p(),r'TOC \t "Figure Caption,1" \h \z')
front_heading('List of tables')
field(p(),r'TOC \t "Table Caption,1" \h \z')
front_heading('List of appendices')
for x in ['A Generative AI declaration','B Retained literature extraction register','C Technical evidence and reproducibility','D Algorithm summaries','E Data dictionary','F Similarity report and submission checks']: p('Appendix '+x)
front_heading('Abbreviations and glossary')
table([['Term','Meaning'],['ACL','Anterior cruciate ligament'],['AUROC','Area under the receiver operating characteristic curve'],['BCE','Binary cross-entropy'],['CRISP-DM','Cross-industry standard process for data mining'],['GPU','Graphics processing unit'],['Grad-CAM','Gradient-weighted class activation mapping'],['MRI','Magnetic resonance imaging'],['PR','Precision–recall'],['ROC','Receiver operating characteristic'],['Study','The grouped MRI input and prediction unit'],['Calibration','Mapping scores toward observed event frequencies on a designated sample'],['External reference','Comparison against a separately sourced reference with stated limits'],['Macro average','Unweighted arithmetic mean across the three targets'],['Threshold','Score boundary used to assign a research finding status']])

for sty in ['Figure Caption','Table Caption']:
    s=doc.styles.add_style(sty,1); s.base_style=doc.styles['Caption']
    s.font.name='Times New Roman'; s.font.size=Pt(10)
    s.paragraph_format.line_spacing=1.15
doc.styles['Table Caption'].paragraph_format.keep_with_next=True
sec=doc.add_section(WD_SECTION_START.NEW_PAGE); number_section(sec,'decimal',1)
text=(OUT/'report_content.md').read_text(encoding='utf-8')
lines=text.splitlines(); i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',x or '-') for x in row): rows.append(row)
            i+=1
        table(rows); continue
    if line.startswith('# '):
        q=doc.add_heading(line[2:],1)
        if line.startswith('# CHAPTER 1'): q.paragraph_format.page_break_before=False
    elif line.startswith('## '): doc.add_heading(line[3:],2)
    elif line.startswith('### '): doc.add_heading(line[4:],3)
    elif line.startswith('@figure '):
        path,caption=line[8:].split(' | ',1)
        q=p(); q.paragraph_format.keep_with_next=True
        q.add_run().add_picture(str(ROOT/path),width=Inches(6.1))
        p(caption,'Figure Caption')
    elif line.startswith('Table '): p(line,'Table Caption')
    else: p(line)
    i+=1

# APA references from verified DOI metadata, with publication-specific corrections.
rs=json.loads((OUT/'qa/csl.json').read_text(encoding='utf-8'))
titles={
'10.1109/cvpr.2016.90':'Deep residual learning for image recognition',
'10.1109/iccv.2017.74':'Grad-CAM: Visual explanations from deep networks via gradient-based localization',
'10.1148/ryai.2020200029':'Checklist for artificial intelligence in medical imaging (CLAIM): A guide for authors and reviewers',
'10.1214/aos/1176344552':'Bootstrap methods: Another look at the jackknife',
'10.1371/journal.pone.0118432':'The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets',
'10.1097/ede.0b013e3181c30fb2':'Assessing the performance of prediction models: A framework for traditional and novel measures',
'10.1038/s41746-022-00592-y':'Machine learning for medical imaging: Methodological failures and recommendations for the future',
'10.1038/s41597-022-01255-z':'fastMRI+, Clinical pathology annotations for knee and brain fully sampled magnetic resonance imaging data'}
titles['10.1136/bmj-2023-078378']='TRIPOD+AI statement: Updated guidance for reporting clinical prediction models that use regression or machine learning methods'
def initials(g):
    return ' '.join('-'.join(y[0]+'.' for y in x.split('-') if y) for x in g.replace('.','').split())
def authors(au):
    a=[x['family']+', '+initials(x.get('given',''))+((', '+x['suffix'].rstrip('.')+'.') if x.get('suffix') else '') for x in au]
    if len(a)>20: return ', '.join(a[:19])+', … '+a[-1]
    if len(a)>1: return ', '.join(a[:-1])+', & '+a[-1]
    return a[0]
refs=[]
for r in rs:
    doi=r['DOI'].lower(); title=titles.get(doi,r['title']); year=r['issued']['date-parts'][0][0]
    container=r['container-title']; volume=r.get('volume',''); issue=r.get('issue',''); pages=r.get('page','')
    if doi=='10.1214/aos/1176344552': pages='1–26'
    if doi=='10.1186/s12916-019-1426-2': pages='Article 195'; issue=''
    if doi.startswith('10.1038/'): pages='Article '+r['article-number']; issue=''
    elif pages and (pages.startswith('e') or doi=='10.1016/j.patter.2023.100804'): pages='Article '+pages
    if doi.startswith('10.1109/'):
        tail=f'In *{container}* (pp. {pages.replace("-","–")}). IEEE.'
    else:
        tail=f'*{container}, {volume}*'+(f'({issue})' if issue else '')+(f', {pages.replace("-","–")}' if pages else '')+'.'
    refs.append((r['author'][0]['family'],f'{authors(r["author"])} ({year}). {title}. {tail} https://doi.org/{r["DOI"]}'))
extra=[
('Adebayo','Adebayo, J., Gilmer, J., Muelly, M., Goodfellow, I., Hardt, M., & Kim, B. (2018). Sanity checks for saliency maps. In *Advances in neural information processing systems* (Vol. 31). https://proceedings.neurips.cc/paper/2018/hash/294a8ed24b1ad22ec2e7efea049b8737-Abstract.html'),
('Guo','Guo, C., Pleiss, G., Sun, Y., & Weinberger, K. Q. (2017). On calibration of modern neural networks. In D. Precup & Y. W. Teh (Eds.), *Proceedings of the 34th international conference on machine learning* (Vol. 70, pp. 1321–1330). PMLR. https://proceedings.mlr.press/v70/guo17a.html'),
('Loshchilov','Loshchilov, I., & Hutter, F. (2019). *Decoupled weight decay regularization*. International Conference on Learning Representations. https://openreview.net/forum?id=Bkg6RiCqY7'),
('Paszke','Paszke, A., Gross, S., Massa, F., Lerer, A., Bradbury, J., Chanan, G., Killeen, T., Lin, Z., Gimelshein, N., Antiga, L., Desmaison, A., Köpf, A., Yang, E., DeVito, Z., Raison, M., Tejani, A., Chilamkurthy, S., Steiner, B., Fang, L., … Chintala, S. (2019). PyTorch: An imperative style, high-performance deep learning library. In *Advances in neural information processing systems* (Vol. 32). https://proceedings.neurips.cc/paper/2019/hash/bdbca288fee7f92f2bfa9f7012727740-Abstract.html'),
('Srivastava','Srivastava, N., Hinton, G., Krizhevsky, A., Sutskever, I., & Salakhutdinov, R. (2014). Dropout: A simple way to prevent neural networks from overfitting. *Journal of Machine Learning Research, 15*, 1929–1958. https://www.jmlr.org/papers/v15/srivastava14a.html'),
('Tan','Tan, M., & Le, Q. (2019). EfficientNet: Rethinking model scaling for convolutional neural networks. In K. Chaudhuri & R. Salakhutdinov (Eds.), *Proceedings of the 36th international conference on machine learning* (Vol. 97, pp. 6105–6114). PMLR. https://proceedings.mlr.press/v97/tan19a.html'),
('Youden','Youden, W. J. (1950). Index for rating diagnostic tests. *Cancer, 3*(1), 32–35. https://doi.org/10.1002/1097-0142(1950)3:1%3C32::AID-CNCR2820030106%3E3.0.CO;2-3'),
('Chapman','Chapman, P., Clinton, J., Kerber, R., Khabaza, T., Reinartz, T., Shearer, C., & Wirth, R. (2000). *CRISP-DM 1.0: Step-by-step data mining guide*. CRISP-DM Consortium. https://nas.uhcl.edu/boetticher/ML_DataMining/CRISPWP-0800.pdf'),
('Midlands','Midlands State University, Department of Information and Marketing Sciences. (n.d.). *Level 2.2 and Level 4.2 data science and informatics project guidelines* (GL-IMS-002) [Unpublished departmental guidelines supplied by the student].')]
refs+=extra
refs.sort(key=lambda x:unicodedata.normalize('NFKD',x[0]).encode('ascii','ignore').decode().lower())
    p('The bibliography contains 21 refereed publications and two additional sources: the CRISP-DM technical guide and the supplied departmental guidelines. DOI metadata and primary publication records were checked during preparation. The departmental document informed the report structure (Midlands State University, Department of Information and Marketing Sciences, n.d.).')
    doc.add_heading('REFERENCES',1)
for _,ref in refs:
    q=p(); q.paragraph_format.left_indent=Inches(.5); q.paragraph_format.first_line_indent=Inches(-.5)
    q.paragraph_format.keep_together=True
    q.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    parts=ref.split('*')
    for j,s in enumerate(parts): q.add_run(s).italic=bool(j%2)
(OUT/'verified_references.txt').write_text('\n\n'.join(x[1].replace('*','') for x in refs),encoding='utf-8')

doc.add_heading('APPENDIX A GENERATIVE AI DECLARATION',1)
p('Generative AI was used. This draft declaration must be reviewed and signed by the student; it is not a substitute for any prescribed institutional form. The applicable categories are related-work sourcing and summarisation, methods and experiment support, code development, data-analysis presentation and interpretation, graphics and formatting, editing, writing, and citation formatting. Assistance also included refinement of the problem and project design. No theorem proving or new mathematical theory is claimed.')
p('Tool: OpenAI ChatGPT/Codex. Uses: assistance with the KneeAssist implementation and debugging, interpretation of saved experiment artifacts, source discovery and bibliographic checks, generation and revision of this report, and document formatting. Published sources and saved project evidence remain the basis for factual claims; AI output is not treated as a scholarly authority.')
p('The student must verify the account of assistance, understand the submitted code and analysis, correct any remaining errors and disclose additional tools or assistance used outside this record. Model claims and source citations should be checked before submission. No statement of unaided authorship is made.')
p('Student name and registration: ____________________')
p('Signature: ____________________    Date: ____________________')

doc.add_heading('APPENDIX B RETAINED LITERATURE EXTRACTION REGISTER',1)
p('This register documents the retained narrative-review sources, not a systematic-review screening flow. Full publication titles, author lists and persistent links are in the reference list. The following entries identify each source’s role and the limit on its use.')
rows=[['Source','Evidence extracted','Limit when applied here'],
['Bien et al. (2018)','Knee MRI three-target classification','Different local architecture and partitions'],
['Štajduhar et al. (2017)','ACL-focused MRI task','Not a three-label normality reference'],
['Zhao et al. (2022)','fastMRI pathology annotations','Local negative references are incomplete'],
['He et al. (2016)','Residual image encoder','General vision result'],
['Tan and Le (2019)','EfficientNet encoder design','No knee-specific superiority guarantee'],
['Srivastava et al. (2014)','Dropout regularisation','No local dropout ablation'],
['Loshchilov and Hutter (2019)','AdamW optimisation','No local optimiser comparison'],
['Paszke et al. (2019)','PyTorch implementation framework','Not evidence of clinical validity'],
['Steyerberg et al. (2010)','Multiple performance dimensions','Clinical value remains unmeasured'],
['Saito and Rehmsmeier (2015)','Precision–recall interpretation','Depends on cohort composition'],
['Youden (1950)','Threshold index','Does not specify clinical error costs'],
['Guo et al. (2017)','Calibration motivation','Local method differs'],
['Efron (1979)','Bootstrap uncertainty','Conditional study resampling only'],
['Kapoor and Narayanan (2023)','Leakage and reproducibility','Audit cannot prove all independence'],
['Varoquaux and Cheplygina (2022)','Evaluation weaknesses','Guidance rather than local evidence'],
['Zech et al. (2018)','Cross-source transfer concern','Chest radiographs rather than knee MRI'],
['Selvaraju et al. (2017)','Grad-CAM explanation method','Not lesion segmentation'],
['Adebayo et al. (2018)','Saliency sanity checks','Checks not yet completed locally'],
['Mongan et al. (2020)','CLAIM reporting guidance','No independent compliance assessment'],
['Collins et al. (2024)','TRIPOD+AI reporting guidance','Not regulatory or clinical approval'],
['Kelly et al. (2019)','Clinical translation challenges','No local outcome study']]
table(rows)

doc.add_heading('APPENDIX C TECHNICAL EVIDENCE AND REPRODUCIBILITY',1)
p('Project root: C:\\Users\\HP\\Desktop\\KneeAssist-AI. Paths below are relative to that root. This report documents existing results; document preparation did not retrain or change the model.')
table([['Artifact','Purpose'],['runs/mri_finetune_02/READINESS_REPORT.md','Release interpretation and workflow verification'],['runs/mri_finetune_02/validation_summary.json','MRNet and external-reference results'],['runs/mri_finetune_02/results/model_comparison.json','Internal selection evidence and limits'],['runs/mri_finetune_02/results/calibration.json','Calibration parameters and thresholds'],['runs/mri_finetune_02/results/leakage_audit.json','Study and series audit'],['runs/mri_finetune_02/results/final/metrics.json','Per-target final metrics'],['runs/mri_finetune_02/results/final/uncertainty.json','Bootstrap uncertainty'],['logs/defence_final_pytest_20261008.log','43 passing automated tests in final regression check'],['runs/model_family_v2/FINAL_DECISION.md','Five-family benchmark and non-promotion decision'],['results/notebook_execution.json','Nine executed notebooks'],['docs/RSNA_ONLINE_TRAINING_STATUS.md','Online pilot status and limitations'],['docs/RSNA_LABEL_REVIEW_PROTOCOL.md','Future review gate; no unverified training labels'],['runs/mri_finetune_02/browser_export_verified.json','Real browser export evidence']])
p('Formal historical baseline checkpoint: runs/mri_finetune_02/models/best_model.pth (EfficientNet-B0; MRNet-only; frozen warm-up followed by calibration). The current dashboard routes DenseNet-121 to general abnormality, ResNet-18 to ACL and Swin-T to meniscus through model_registry/active_models.json before scoring; it is not an ensemble or patient-specific model selection.')
p('Checkpoint SHA-256: b60293f6c7130fd646c81b6d758584faf3e17b2689503b5278ad711156d546da')
p('Before rerunning an experiment, preserve this run and its configuration, confirm the dataset manifests, and use a new run directory. The current evaluation cohorts must not be re-described as unseen data. A refreshed experiment should state its own split and selection rules.')

doc.add_heading('APPENDIX D ALGORITHM SUMMARIES',1)
doc.add_heading('D.1 Training and selection',2)
for s in ['1. Load study manifests and target labels; retain the stated dataset roles.', '2. Verify study separation and exact-series checks; record unavailable patient linkage.', '3. Prepare within-volume normalisation and sampled slices using the fixed configuration.', '4. Fit the two primary candidate architectures on the fitting partition; monitor internal tuning AUROC.', '5. Retain the separate five-family portfolio as supplementary target-routing evidence; update a route only under the locked validation policy, never because a model gives a higher individual patient score.', '6. Select the formal baseline checkpoint using internal macro AUROC; maintain live target routes only through the locked registry policy.', '7. Fit positive-slope calibration and research thresholds on the reserved calibration studies.', '8. Evaluate the frozen release on the stated cohorts and save all results, including failures.']: p(s)
doc.add_heading('D.2 Inference and presentation',2)
for s in ['1. Accept supported MRI arrays or a ZIP containing them and a pseudonymous case reference.', '2. Validate the supported file structure; report readable errors for invalid inputs.', '3. Apply the same preprocessing and load the registered target-specific checkpoint for each reported finding.', '4. Aggregate slice and plane features, calculate logits and apply saved calibration.', '5. Compare each score with its own saved threshold and generate requested attention.', '6. Present scores, research finding status, slice views and the clinical-review notice.', '7. Export the case summary with model provenance or clear the case.']: p(s)

doc.add_heading('APPENDIX E DATA DICTIONARY',1)
table([['Concept','Representation','Interpretation'],['Study reference','Identifier','Grouping key; not proof of unique patient'],['Axial stack','Numeric slice array','One available MRI plane'],['Coronal stack','Numeric slice array','One available MRI plane'],['Sagittal stack','Numeric slice array','One available MRI plane'],['abnormal','Binary reference / score','General abnormality target'],['acl','Binary reference / score','ACL tear target'],['meniscus','Binary reference / score','Meniscal tear target'],['Availability mask','Plane indicators','Records supplied planes'],['Threshold','Target-specific number','Research status boundary'],['Checkpoint digest','SHA-256 string','Identity of model file'],['Attention map','Class-related image overlay','Model attention; not confirmed lesion']])
p('This is a conceptual dictionary of the variables used in the report. It does not imply a universal DICOM schema or that all three references exist in every external source.')

doc.add_heading('APPENDIX F SIMILARITY REPORT AND SUBMISSION CHECKS',1)
p('A Turnitin or other institutional similarity report has not been supplied or generated. No similarity percentage or plagiarism clearance is claimed. Attach the authentic institutional report here after the student completes the required submission process.')
p('Before submission, complete the cover identity fields and campus, confirm the project level, review the declaration, obtain genuine supervisor approval, add any personal dedication or acknowledgements, and attach the required similarity report. Update the Word contents and caption lists after any final edits. The descriptive file name should be replaced by the institution’s required name, surname and registration-number format once those details are confirmed.')
p('The report follows the supplied five-chapter structure, uses exactly five specific objectives beginning with “To”, and includes an explicit AI declaration. The student remains responsible for confirming department-specific page-length and submission requirements with the supervisor.')
settings=doc.settings.element
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
# Remove theme-font and decorative-border defaults that otherwise override the
# explicit university typography when Microsoft Word renders the document.
for st in doc.styles:
    if hasattr(st,'font'):
        st.font.name='Times New Roman'
        rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        if rf is not None:
            for attr in list(rf.attrib):
                if attr.endswith('Theme'): del rf.attrib[attr]
            for attr in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+attr),'Times New Roman')
    for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
for q in doc.paragraphs:
    for border in list(q._p.iter(qn('w:pBdr'))): border.getparent().remove(border)
dest=OUT/'KneeAssist_AI_Five_Objective_Project_Report.docx'
doc.save(dest)
print(json.dumps({'docx':str(dest),'abstract_words':len(abstract.split()),'body_words':len(text.split()),'references':len(refs),'refereed_sources':len(refs)-2},indent=2))
